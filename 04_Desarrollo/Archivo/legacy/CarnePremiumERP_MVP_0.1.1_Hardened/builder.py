from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import json
import shutil
import tempfile
import uuid

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from config import MASTER_FILE, SHEETS

MAX_ROWS = 500
BUILD_VERSION = "0.1.1"
BUILD_NAME = "MVP Clientes + Productos + Ventas"
HEADER_FILL = "7A3E13"
HEADER_FONT = "FFFFFF"
STATE_FILE = MASTER_FILE.parent / ".build_state.json"


def sha256_file(path: Path) -> str:
    sha = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            sha.update(chunk)
    return sha.hexdigest()


def ensure_workbook(path: Path):
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def ensure_sheet(wb, name: str):
    if name not in wb.sheetnames:
        wb.create_sheet(name)
    return wb[name]


def style_header(ws, row: int = 1):
    for cell in ws[row]:
        if cell.value is not None:
            cell.font = Font(bold=True, color=HEADER_FONT)
            cell.fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def clear_validations(ws):
    ws.data_validations.dataValidation = []


def add_list_validation(ws, target_range: str, formula1: str | None = None,
                        values: list[str] | None = None):
    if formula1 is None and values is None:
        raise ValueError("Debe indicar formula1 o values.")
    dv = DataValidation(type="list", allow_blank=True)
    dv.formula1 = formula1 if formula1 else '"' + ",".join(values or []) + '"'
    dv.errorTitle = "Valor no válido"
    dv.error = "Seleccione un valor permitido."
    dv.promptTitle = "Carne Premium ERP"
    dv.prompt = "Seleccione un valor de la lista."
    ws.add_data_validation(dv)
    dv.add(target_range)
    return dv


def ensure_table(ws, table_name: str, ref: str):
    if table_name in [t.name for t in ws.tables.values()]:
        del ws.tables[table_name]
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def set_headers(ws, headers):
    for col, header in enumerate(headers, start=1):
        ws.cell(1, col).value = header
    style_header(ws)


def setup_parameters(wb):
    ws = ensure_sheet(wb, "01_PARAMETROS")
    set_headers(ws, ["Parametro", "Valor"])
    defaults = [
        ("Empresa", "Carne Premium"),
        ("Moneda", "ARS"),
        ("TipoPrecioDefault", "Minorista"),
        ("MedioPagoDefault", "Efectivo"),
        ("IVA", 21),
    ]
    for row, (key, value) in enumerate(defaults, 2):
        if not ws.cell(row, 1).value:
            ws.cell(row, 1).value = key
        if ws.cell(row, 2).value is None:
            ws.cell(row, 2).value = value
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 24


def setup_catalogs(wb):
    ws = ensure_sheet(wb, "02_CATALOGOS")
    set_headers(ws, ["TiposPrecio", "TiposCliente", "MediosPago", "Estados"])
    values = [
        ["Minorista", "Minorista", "Efectivo", "Activo"],
        ["Mayorista", "Mayorista", "Transferencia", "Inactivo"],
        ["Especial", "Especial", "Cuenta Corriente", None],
        [None, None, "Tarjeta", None],
        [None, None, "Otro", None],
    ]
    for r, row_values in enumerate(values, 2):
        for c, value in enumerate(row_values, 1):
            if ws.cell(r, c).value is None:
                ws.cell(r, c).value = value
    for col, width in {"A":20,"B":20,"C":22,"D":16}.items():
        ws.column_dimensions[col].width = width


def setup_clients(wb):
    ws = ensure_sheet(wb, "03_CLIENTES")
    clear_validations(ws)
    set_headers(ws, ["ClienteID","Nombre","Telefono","Direccion","TipoCliente","Estado","Observaciones"])
    for row in range(2, MAX_ROWS + 1):
        ws.cell(row,1).value = f'=IF(B{row}="","","CLI"&TEXT(ROW()-1,"000000"))'
        if not ws.cell(row,6).value:
            ws.cell(row,6).value = "Activo"
    add_list_validation(ws, "E2:E501", formula1="'02_CATALOGOS'!$B$2:$B$4")
    add_list_validation(ws, "F2:F501", formula1="'02_CATALOGOS'!$D$2:$D$3")
    ensure_table(ws, "ClientesTable", "A1:G501")
    for col, width in {"A":15,"B":28,"C":18,"D":30,"E":18,"F":14,"G":30}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes.freeze_rows(1)


def setup_products(wb):
    ws = ensure_sheet(wb, "04_PRODUCTOS")
    clear_validations(ws)
    set_headers(ws, ["ProductoID","Producto","Categoria","Costo","PrecioMinorista","PrecioMayorista","PrecioEspecial","Estado","Observaciones"])
    for row in range(2, MAX_ROWS + 1):
        ws.cell(row,1).value = f'=IF(B{row}="","","PRO"&TEXT(ROW()-1,"000000"))'
        if not ws.cell(row,8).value:
            ws.cell(row,8).value = "Activo"
    add_list_validation(ws, "H2:H501", formula1="'02_CATALOGOS'!$D$2:$D$3")
    ws["D2:G501"].number_format = '$ #,##0.00'
    ensure_table(ws, "ProductosTable", "A1:I501")
    for col, width in {"A":15,"B":28,"C":20,"D":16,"E":18,"F":18,"G":18,"H":14,"I":30}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes.freeze_rows(1)


def setup_sales(wb):
    ws = ensure_sheet(wb, "05_VENTAS")
    clear_validations(ws)
    set_headers(ws, ["VentaID","Fecha","Hora","ClienteID","TipoPrecio","ProductoID","Kg","PrecioKg","Importe","Ganancia$","Ganancia%","MedioPago","Vendedor","Observaciones"])
    for row in range(2, MAX_ROWS + 1):
        ws.cell(row,1).value = f'=IF(OR(D{row}="",F{row}=""),"","VEN"&TEXT(ROW()-1,"000000"))'
        ws.cell(row,2).value = f'=IF(A{row}="","",TODAY())'
        ws.cell(row,3).value = f'=IF(A{row}="","",NOW())'
    add_list_validation(ws, "D2:D501", formula1="'03_CLIENTES'!$A$2:$A$501")
    add_list_validation(ws, "E2:E501", formula1="'02_CATALOGOS'!$A$2:$A$4")
    add_list_validation(ws, "F2:F501", formula1="'04_PRODUCTOS'!$A$2:$A$501")
    add_list_validation(ws, "L2:L501", formula1="'02_CATALOGOS'!$C$2:$C$6")
    for row in range(2, MAX_ROWS + 1):
        ws.cell(row,8).value = (
            f'=IFERROR(IF(E{row}="Mayorista",IF(VLOOKUP(F{row},\'04_PRODUCTOS\'!$A$2:$G$501,6,FALSE)="",VLOOKUP(F{row},\'04_PRODUCTOS\'!$A$2:$G$501,5,FALSE),VLOOKUP(F{row},\'04_PRODUCTOS\'!$A$2:$G$501,6,FALSE)),IF(E{row}="Especial",IF(VLOOKUP(F{row},\'04_PRODUCTOS\'!$A$2:$G$501,7,FALSE)="",VLOOKUP(F{row},\'04_PRODUCTOS\'!$A$2:$G$501,5,FALSE),VLOOKUP(F{row},\'04_PRODUCTOS\'!$A$2:$G$501,7,FALSE)),VLOOKUP(F{row},\'04_PRODUCTOS\'!$A$2:$G$501,5,FALSE))),"")'
        )
        ws.cell(row,9).value = f'=IF(AND(G{row}<>"",H{row}<>""),G{row}*H{row},"")'
        ws.cell(row,10).value = f'=IFERROR(IF(VLOOKUP(F{row},\'04_PRODUCTOS\'!$A$2:$D$501,4,FALSE)="","",I{row}-(G{row}*VLOOKUP(F{row},\'04_PRODUCTOS\'!$A$2:$D$501,4,FALSE))),"")'
        ws.cell(row,11).value = f'=IF(OR(I{row}="",J{row}=""),"",IFERROR(J{row}/I{row},""))'
    ws["G2:G501"].number_format = "0.000"
    ws["H2:J501"].number_format = '$ #,##0.00'
    ws["K2:K501"].number_format = "0.0%"
    ws["B2:B501"].number_format = "dd/mm/yyyy"
    ws["C2:C501"].number_format = "hh:mm:ss"
    ensure_table(ws, "VentasTable", "A1:N501")
    for col, width in {"A":15,"B":12,"C":10,"D":15,"E":15,"F":15,"G":10,"H":16,"I":16,"J":16,"K":12,"L":20,"M":20,"N":30}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes.freeze_rows(1)


def setup_home(wb):
    ws = ensure_sheet(wb, "00_INICIO")
    ws.unmerge_cells("A1:B1") if "A1:B1" in [str(rng) for rng in ws.merged_cells.ranges] else None
    ws.merge_cells("A1:B1")
    ws["A1"] = "CARNE PREMIUM ERP"
    ws["A2"] = "MVP 0.1.1"
    rows = [
        ("Flujo de prueba", "Clientes → Productos → Ventas"),
        ("Paso 1", "Cargar cliente en 03_CLIENTES"),
        ("Paso 2", "Cargar producto y precios en 04_PRODUCTOS"),
        ("Paso 3", "Registrar venta en 05_VENTAS"),
        ("Paso 4", "Verificar precio, importe y ganancia"),
        ("Estado", "MVP 0.1.1"),
        ("Compatibilidad", "Excel 2019 + Google Sheets"),
    ]
    for r, (key, value) in enumerate(rows, 4):
        ws.cell(r,1).value = key
        ws.cell(r,2).value = value
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=16)
    ws["A1"].fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 52


def setup_version(wb, build_token):
    ws = ensure_sheet(wb, "12_VERSION")
    set_headers(ws, ["Versión","Build","Estado","BuildToken"])
    ws["B1"] = BUILD_VERSION
    ws["A2"] = "Build"; ws["B2"] = BUILD_NAME
    ws["A3"] = "Estado"; ws["B3"] = "Generado por builder"
    ws["A4"] = "BuildToken"; ws["B4"] = build_token
    style_header(ws)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 60


def setup_log(wb, build_token):
    ws = ensure_sheet(wb, "13_LOG")
    set_headers(ws, ["Fecha","Versión","Acción","Detalle"])
    tokens = {str(row[3]) for row in ws.iter_rows(min_row=2, values_only=True) if len(row) >= 4 and row[3]}
    if build_token not in tokens:
        r = ws.max_row + 1
        ws.cell(r,1).value = datetime.now()
        ws.cell(r,2).value = BUILD_VERSION
        ws.cell(r,3).value = "BUILD"
        ws.cell(r,4).value = build_token
    for col, width in {"A":20,"B":14,"C":14,"D":50}.items():
        ws.column_dimensions[col].width = width


def validate_before_save(wb, expected_token):
    required = {
        "03_CLIENTES": ["ClienteID","Nombre","Telefono","Direccion","TipoCliente","Estado","Observaciones"],
        "04_PRODUCTOS": ["ProductoID","Producto","Categoria","Costo","PrecioMinorista","PrecioMayorista","PrecioEspecial","Estado","Observaciones"],
        "05_VENTAS": ["VentaID","Fecha","Hora","ClienteID","TipoPrecio","ProductoID","Kg","PrecioKg","Importe","Ganancia$","Ganancia%","MedioPago","Vendedor","Observaciones"],
        "12_VERSION": ["Versión","Build","Estado","BuildToken"],
    }
    for sheet, headers in required.items():
        if sheet not in wb.sheetnames:
            raise RuntimeError(f"Falta hoja obligatoria: {sheet}")
        actual = [c.value for c in wb[sheet][1]]
        missing = [h for h in headers if h not in actual]
        if missing:
            raise RuntimeError(f"{sheet}: faltan encabezados {missing}")
    version = wb["12_VERSION"]
    if version["B1"].value != BUILD_VERSION:
        raise RuntimeError("La versión del libro no coincide con el builder.")
    if version["B4"].value != expected_token:
        raise RuntimeError("El BuildToken no coincide.")
    sales = wb["05_VENTAS"]
    for cell in ["A2","B2","C2","H2","I2","J2","K2"]:
        if not isinstance(sales[cell].value, str) or not sales[cell].value.startswith("="):
            raise RuntimeError(f"05_VENTAS!{cell} no contiene una fórmula.")
    if len(wb["03_CLIENTES"].data_validations.dataValidation) < 2:
        raise RuntimeError("03_CLIENTES no tiene sus validaciones esperadas.")
    if len(wb["04_PRODUCTOS"].data_validations.dataValidation) < 1:
        raise RuntimeError("04_PRODUCTOS no tiene su validación esperada.")
    if len(wb["05_VENTAS"].data_validations.dataValidation) < 4:
        raise RuntimeError("05_VENTAS no tiene las cuatro listas esperadas.")


def atomic_save(wb):
    build_token = f"{BUILD_VERSION}-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:8]}"
    setup_version(wb, build_token)
    setup_log(wb, build_token)
    validate_before_save(wb, build_token)

    MASTER_FILE.parent.mkdir(parents=True, exist_ok=True)
    backup_dir = MASTER_FILE.parent.parent / "02_Backups"
    backup_dir.mkdir(parents=True, exist_ok=True)

    fd, tmp_name = tempfile.mkstemp(prefix="CarnePremiumERP_", suffix=".xlsx", dir=str(MASTER_FILE.parent))
    Path(tmp_name).unlink(missing_ok=True)
    try:
        wb.save(tmp_name)
        check = load_workbook(tmp_name, data_only=False)
        try:
            validate_before_save(check, build_token)
        finally:
            check.close()

        if MASTER_FILE.exists():
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            shutil.copy2(MASTER_FILE, backup_dir / f"CarnePremiumERP_before_{stamp}.xlsx")

        Path(tmp_name).replace(MASTER_FILE)
        digest = sha256_file(MASTER_FILE)
        STATE_FILE.write_text(json.dumps({
            "version": BUILD_VERSION,
            "build": BUILD_NAME,
            "build_token": build_token,
            "master_file": str(MASTER_FILE),
            "sha256": digest,
            "saved_at": datetime.now().isoformat(timespec="seconds"),
        }, indent=2, ensure_ascii=False), encoding="utf-8")
        return build_token
    except PermissionError as exc:
        raise PermissionError("No se pudo reemplazar CarnePremiumERP.xlsx. Cerrá Excel y cualquier otro programa que tenga abierto el archivo y volvé a ejecutar el builder.") from exc
    finally:
        Path(tmp_name).unlink(missing_ok=True)


def main():
    print("=" * 50)
    print("CARNE PREMIUM ERP - MVP 0.1.1")
    print("=" * 50)
    wb = ensure_workbook(MASTER_FILE)
    print("[1/4] Parámetros y catálogos...")
    setup_parameters(wb)
    setup_catalogs(wb)
    print("[2/4] Clientes...")
    setup_clients(wb)
    print("[3/4] Productos y Ventas...")
    setup_products(wb)
    setup_sales(wb)
    setup_home(wb)
    print("[4/4] Validando y guardando...")
    token = atomic_save(wb)
    print("\nOK: ERP generado correctamente.")
    print(f"Versión: {BUILD_VERSION}")
    print(f"BuildToken: {token}")
    print(f"Archivo: {MASTER_FILE}")
    print(f"Estado: {STATE_FILE}")

if __name__ == "__main__":
    main()
