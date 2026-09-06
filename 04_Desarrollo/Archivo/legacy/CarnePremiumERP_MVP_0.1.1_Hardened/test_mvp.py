from __future__ import annotations

from pathlib import Path
import hashlib
import json
import sys
from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
SCRIPTS_DIR = SCRIPT_DIR.parent / "Scripts"
sys.path.insert(0, str(SCRIPTS_DIR))
from config import MASTER_FILE, SHEETS

BUILD_VERSION = "0.1.1"
BUILD_NAME = "MVP Clientes + Productos + Ventas"
STATE_FILE = MASTER_FILE.parent / ".build_state.json"


def sha256_file(path: Path) -> str:
    sha = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            sha.update(chunk)
    return sha.hexdigest()


def fail(message: str):
    print(f"ERROR: {message}")
    raise SystemExit(1)


def main():
    if not MASTER_FILE.exists():
        fail(f"No existe el archivo maestro: {MASTER_FILE}")
    if not STATE_FILE.exists():
        fail("No existe .build_state.json. El builder probablemente no terminó correctamente.")

    state = json.loads(STATE_FILE.read_text(encoding="utf-8"))
    actual_hash = sha256_file(MASTER_FILE)
    if actual_hash != state.get("sha256"):
        fail("EL ARCHIVO MAESTRO NO COINCIDE CON EL ÚLTIMO BUILD. Puede haber quedado una versión vieja porque el builder falló.")
    if state.get("version") != BUILD_VERSION:
        fail(f"Se esperaba versión {BUILD_VERSION}, pero el estado indica {state.get('version')}.")

    wb = load_workbook(MASTER_FILE, data_only=False)
    try:
        missing_sheets = [s for s in SHEETS if s not in wb.sheetnames]
        if missing_sheets:
            fail(f"Faltan hojas obligatorias: {missing_sheets}")

        required_headers = {
            "03_CLIENTES": ["ClienteID","Nombre","Telefono","Direccion","TipoCliente","Estado","Observaciones"],
            "04_PRODUCTOS": ["ProductoID","Producto","Categoria","Costo","PrecioMinorista","PrecioMayorista","PrecioEspecial","Estado","Observaciones"],
            "05_VENTAS": ["VentaID","Fecha","Hora","ClienteID","TipoPrecio","ProductoID","Kg","PrecioKg","Importe","Ganancia$","Ganancia%","MedioPago","Vendedor","Observaciones"],
        }
        for sheet, expected in required_headers.items():
            actual = [c.value for c in wb[sheet][1]]
            missing = [h for h in expected if h not in actual]
            if missing:
                fail(f"{sheet}: faltan encabezados {missing}")

        version = wb["12_VERSION"]
        if version["B1"].value != BUILD_VERSION:
            fail(f"El Excel no está en {BUILD_VERSION}; está en {version['B1'].value}.")
        if version["B2"].value != BUILD_NAME:
            fail("El nombre de build no coincide.")
        if version["B4"].value != state.get("build_token"):
            fail("EL EXCEL NO CORRESPONDE AL ÚLTIMO BUILD. El BuildToken no coincide.")

        sales = wb["05_VENTAS"]
        for cell in ["A2","B2","C2","H2","I2","J2","K2"]:
            value = sales[cell].value
            if not isinstance(value, str) or not value.startswith("="):
                fail(f"05_VENTAS!{cell} no contiene la fórmula esperada.")

        if len(wb["03_CLIENTES"].data_validations.dataValidation) < 2:
            fail("03_CLIENTES no tiene las validaciones esperadas.")
        if len(wb["04_PRODUCTOS"].data_validations.dataValidation) < 1:
            fail("04_PRODUCTOS no tiene la validación esperada.")
        if len(wb["05_VENTAS"].data_validations.dataValidation) < 4:
            fail("05_VENTAS no tiene las cuatro validaciones esperadas.")

        print("=" * 60)
        print("OK: MVP 0.1.1 VALIDADO")
        print("=" * 60)
        print(f"Archivo: {MASTER_FILE}")
        print(f"Versión: {version['B1'].value}")
        print(f"BuildToken: {version['B4'].value}")
        print("Clientes: estructura + validaciones OK")
        print("Productos: estructura + validaciones OK")
        print("Ventas: estructura + fórmulas + validaciones OK")
        print("Protección contra archivo maestro viejo: OK")
        print("Integridad SHA-256 del último build: OK")
    finally:
        wb.close()

if __name__ == "__main__":
    main()
