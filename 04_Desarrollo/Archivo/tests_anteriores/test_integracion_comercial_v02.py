from __future__ import annotations

from datetime import datetime

import sys
from pathlib import Path

from openpyxl import Workbook


TEST_DIR = Path(
    __file__
).resolve().parent

SCRIPTS_DIR = (
    TEST_DIR.parent
    / "Scripts"
)

sys.path.insert(
    0,
    str(SCRIPTS_DIR)
)


from ventas_integradas import (
    procesar_venta
)


# ============================================================
# WORKBOOK DE PRUEBA
# ============================================================

def crear_workbook():

    wb = Workbook()

    wb.remove(
        wb.active
    )

    # ========================================================
    # CLIENTES
    # ========================================================

    ws = wb.create_sheet(
        "03_CLIENTES"
    )

    ws.append([
        "ClienteID",
        "Nombre",
        "Telefono",
        "Direccion",
        "TipoCliente",
        "Estado",
        "Observaciones",
    ])

    ws.append([
        "CLI001",
        "Cliente Minorista",
        "",
        "",
        "Minorista",
        "Activo",
        "",
    ])

    ws.append([
        "CLI002",
        "Cliente Mayorista",
        "",
        "",
        "Mayorista",
        "Activo",
        "",
    ])

    ws.append([
        "CLI003",
        "Cliente Especial",
        "",
        "",
        "Minorista",
        "Activo",
        "",
    ])

    ws.append([
        "CLI004",
        "Cliente Oferta",
        "",
        "",
        "Minorista",
        "Activo",
        "",
    ])

    # ========================================================
    # PRODUCTOS
    # ========================================================

    ws = wb.create_sheet(
        "04_PRODUCTOS"
    )

    ws.append([
        "ProductoID",
        "Producto",
        "Categoria",
        "Costo",
        "PrecioMinorista",
        "PrecioMayorista",
        "PrecioEspecial",
        "Estado",
        "Observaciones",
    ])

    ws.append([
        "PRO001",
        "Asado",
        "Vacuno",
        10000,
        16000,
        14500,
        None,
        "Activo",
        "",
    ])

    # ========================================================
    # PRECIOS ESPECIALES
    # ========================================================

    ws = wb.create_sheet(
        "22_PRECIOS"
    )

    ws.append([
        "PrecioID",
        "ClienteID",
        "Cliente",
        "ProductoID",
        "Producto",
        "TipoPrecio",
        "Precio",
        "Descuento%",
        "FechaInicio",
        "FechaFin",
        "Estado",
        "Prioridad",
        "Motivo",
    ])

    ws.append([
        "PRE001",
        "CLI003",
        "Cliente Especial",
        "PRO001",
        "Asado",
        "Especial",
        14000,
        None,
        datetime(
            2026,
            9,
            1
        ),
        datetime(
            2026,
            9,
            30
        ),
        "Activo",
        10,
        "Precio especial cliente",
    ])

    # ========================================================
    # OFERTAS
    # ========================================================

    ws = wb.create_sheet(
        "23_OFERTAS"
    )

    ws.append([
        "OfertaID",
        "ProductoID",
        "Producto",
        "NombreOferta",
        "PrecioOferta",
        "Descuento%",
        "FechaInicio",
        "FechaFin",
        "MinKg",
        "Estado",
        "Prioridad",
    ])

    ws.append([
        "OF001",
        "PRO001",
        "Asado",
        "Oferta Septiembre",
        12000,
        None,
        datetime(
            2026,
            9,
            1
        ),
        datetime(
            2026,
            9,
            30
        ),
        2,
        "Activo",
        100,
    ])

    # ========================================================
    # VENTAS
    # ========================================================

    ws = wb.create_sheet(
        "05_VENTAS"
    )

    ws.append([
        "VentaID",
        "Fecha",
        "Hora",
        "Cliente",
        "TipoPrecio",
        "Producto",
        "Kg",
        "PrecioKg",
        "Importe",
        "Ganancia$",
        "Ganancia%",
        "MedioPago",
        "Vendedor",
        "Observaciones",
        "_ClienteID",
        "_ProductoID",
        "EstadoVenta",
        "TipoAplicado",
        "MotivoPrecio",
        "OfertaID",
        "OfertaNombre",
        "PrecioEspecialID",
    ])

    # Cuatro filas de laboratorio.

    # Fila 2: minorista
    ws.append([
        None,
        datetime(2026, 8, 20),
        None,
        "Cliente Minorista",
        "Minorista",
        "Asado",
        2,
    ])

    # Fila 3: mayorista
    ws.append([
        None,
        datetime(2026, 8, 20),
        None,
        "Cliente Mayorista",
        "Mayorista",
        "Asado",
        2,
    ])

    # Fila 4: especial
    ws.append([
        None,
        datetime(2026, 9, 10),
        None,
        "Cliente Especial",
        "Especial",
        "Asado",
        2,
    ])

    # Fila 5: oferta
    ws.append([
        None,
        datetime(2026, 9, 15),
        None,
        "Cliente Oferta",
        "Minorista",
        "Asado",
        2,
    ])

    return wb


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 70)
    print(
        "INTEGRACION COMERCIAL v0.2"
    )
    print("=" * 70)

    wb = crear_workbook()

    try:

        # ====================================================
        # MINORISTA
        # ====================================================

        print()
        print(
            "CASO 1 - MINORISTA"
        )

        r = procesar_venta(
            wb,
            2
        )

        print(r)

        assert r["estado"] == "OK"
        assert r["precio_kg"] == 16000
        assert r["importe"] == 32000
        assert r["ganancia"] == 12000
        assert r["tipo_aplicado"] == "Minorista"

        # ====================================================
        # MAYORISTA
        # ====================================================

        print()
        print(
            "CASO 2 - MAYORISTA"
        )

        r = procesar_venta(
            wb,
            3
        )

        print(r)

        assert r["estado"] == "OK"
        assert r["precio_kg"] == 14500
        assert r["importe"] == 29000
        assert r["ganancia"] == 9000
        assert r["tipo_aplicado"] == "Mayorista"

        # ====================================================
        # ESPECIAL
        # ====================================================

        print()
        print(
            "CASO 3 - ESPECIAL"
        )

        r = procesar_venta(
            wb,
            4
        )

        print(r)

        assert r["estado"] == "OK"
        assert r["precio_kg"] == 14000
        assert r["importe"] == 28000
        assert r["ganancia"] == 8000
        assert r["tipo_aplicado"] == "Especial"
        assert r["precio_especial_id"] == "PRE001"

        # ====================================================
        # OFERTA
        # ====================================================

        print()
        print(
            "CASO 4 - OFERTA"
        )

        r = procesar_venta(
            wb,
            5
        )

        print(r)

        assert r["estado"] == "OK"
        assert r["precio_kg"] == 12000
        assert r["importe"] == 24000
        assert r["ganancia"] == 4000
        assert r["tipo_aplicado"] == "Oferta"
        assert r["oferta_id"] == "OF001"
        assert r["oferta_nombre"] == "Oferta Septiembre"

        # ====================================================
        # VERIFICAR COLUMNA EXCEL
        # ====================================================

        print()
        print(
            "VERIFICACION 05_VENTAS"
        )

        ws = wb[
            "05_VENTAS"
        ]

        # Fila 5
        assert ws["H5"].value == 12000
        assert ws["I5"].value == 24000
        assert ws["J5"].value == 4000
        assert ws["Q5"].value == "OK"
        assert ws["R5"].value == "Oferta"
        assert ws["S5"].value == "Oferta Septiembre"
        assert ws["T5"].value == "OF001"
        assert ws["U5"].value == "Oferta Septiembre"

        print(
            "PrecioKg: OK"
        )

        print(
            "Importe: OK"
        )

        print(
            "Ganancia: OK"
        )

        print(
            "EstadoVenta: OK"
        )

        print(
            "TipoAplicado: OK"
        )

        print(
            "MotivoPrecio: OK"
        )

        print(
            "OfertaID: OK"
        )

        print(
            "OfertaNombre: OK"
        )

        # ====================================================
        # RESULTADO
        # ====================================================

        print()
        print("=" * 70)
        print(
            "OK: INTEGRACION COMERCIAL v0.2 SUPERADA"
        )
        print("=" * 70)

        print()
        print(
            "Motor → Ventas: OK"
        )

        print(
            "Minorista: OK"
        )

        print(
            "Mayorista: OK"
        )

        print(
            "Especial: OK"
        )

        print(
            "Oferta: OK"
        )

        print(
            "Trazabilidad: OK"
        )

        print()
        print(
            "El Excel maestro NO fue modificado."
        )

    finally:

        wb.close()


if __name__ == "__main__":
    main()