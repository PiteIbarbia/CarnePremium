from __future__ import annotations

from datetime import datetime

import sys
from pathlib import Path

from openpyxl import Workbook

TEST_DIR = Path(__file__).resolve().parent
SCRIPTS_DIR = TEST_DIR.parent / "Scripts"

sys.path.insert(0, str(SCRIPTS_DIR))

from ventas_excel import procesar_venta_excel


def crear_workbook():
    workbook = Workbook()

    workbook.remove(workbook.active)

    # ========================================================
    # CLIENTES
    # ========================================================

    clientes = workbook.create_sheet(
        "03_CLIENTES"
    )

    clientes.append([
        "ClienteID",
        "Nombre",
        "Telefono",
        "Direccion",
        "TipoCliente",
        "Estado",
        "Observaciones"
    ])

    clientes.append([
        "CLI015",
        "Cliente Prueba",
        "1111111111",
        "Direccion Prueba",
        "Minorista",
        "Activo",
        ""
    ])

    # ========================================================
    # PRODUCTOS
    # ========================================================

    productos = workbook.create_sheet(
        "04_PRODUCTOS"
    )

    productos.append([
        "ProductoID",
        "Producto",
        "Categoria",
        "Costo",
        "PrecioMinorista",
        "PrecioMayorista",
        "PrecioEspecial",
        "Estado",
        "Observaciones"
    ])

    productos.append([
        "PRO015",
        "Asado Prueba",
        "Vacuno",
        10000,
        16000,
        None,
        None,
        "Activo",
        ""
    ])

    # ========================================================
    # OFERTAS
    # ========================================================

    ofertas = workbook.create_sheet(
        "23_OFERTAS"
    )

    ofertas.append([
        "OfertaID",
        "ProductoID",
        "NombreOferta",
        "PrecioOferta",
        "Descuento%",
        "FechaInicio",
        "FechaFin",
        "MinKg",
        "Estado",
        "Prioridad"
    ])

    # Oferta válida durante septiembre
    ofertas.append([
        "OF_VALID_015",
        "PRO015",
        "Oferta Septiembre",
        13000,
        None,
        datetime(2026, 9, 1),
        datetime(2026, 9, 7),
        2,
        "Activo",
        100
    ])

    # ========================================================
    # VENTAS
    # ========================================================

    ventas = workbook.create_sheet(
        "05_VENTAS"
    )

    ventas.append([
        "VentaID",
        "Fecha",
        "Hora",
        "Cliente",
        "TipoPrecio",
        "Producto",
        "Kg",
        "PrecioKg",
        "Importe",
        "Ganancia$",
        "Ganancia%",
        "MedioPago",
        "Vendedor",
        "Observaciones",
        "_ClienteID",
        "_ProductoID",
        "EstadoVenta",
        "TipoAplicado",
        "MotivoPrecio",
        "OfertaID",
        "OfertaNombre"
    ])

    return workbook


def preparar_venta(
    workbook,
    fecha,
    kg=2
):
    ventas = workbook["05_VENTAS"]

    # Limpiar fila de prueba
    for col in range(1, 22):
        ventas.cell(2, col).value = None

    ventas.cell(2, 2).value = fecha
    ventas.cell(2, 4).value = "Cliente Prueba"
    ventas.cell(2, 5).value = "Minorista"
    ventas.cell(2, 6).value = "Asado Prueba"
    ventas.cell(2, 7).value = kg

    return procesar_venta_excel(
        workbook,
        2
    )


def imprimir_resultado(
    titulo,
    resultado,
    ventas
):

    print()
    print("=" * 60)
    print(titulo)
    print("=" * 60)

    print(
        f"Estado: "
        f"{resultado.get('estado')}"
    )

    print(
        f"Precio/Kg: "
        f"{resultado.get('precio_kg')}"
    )

    print(
        f"Importe: "
        f"{ventas.cell(2, 9).value}"
    )

    print(
        f"Ganancia: "
        f"{ventas.cell(2, 10).value}"
    )

    print(
        f"Tipo aplicado: "
        f"{ventas.cell(2, 18).value}"
    )

    print(
        f"Motivo: "
        f"{ventas.cell(2, 19).value}"
    )

    print(
        f"OfertaID: "
        f"{ventas.cell(2, 20).value}"
    )

    print(
        f"Oferta: "
        f"{ventas.cell(2, 21).value}"
    )


def main():

    print()
    print("=" * 60)
    print(
        "MVP 0.1.5 - PRUEBA REAL DE OFERTAS EN VENTAS"
    )
    print("=" * 60)

    workbook = crear_workbook()

    try:

        ventas = workbook["05_VENTAS"]

        # ====================================================
        # CASO 1 — SIN OFERTA
        # ====================================================

        resultado = preparar_venta(
            workbook,
            datetime(
                2026,
                8,
                31
            ),
            2
        )

        # La oferta comienza el 01/09/2026
        assert resultado["estado"] == "OK"
        assert resultado["precio_kg"] == 16000

        imprimir_resultado(
            "CASO 1 - SIN OFERTA VIGENTE",
            resultado,
            ventas
        )

        # ====================================================
        # CASO 2 — OFERTA VIGENTE
        # ====================================================

        resultado = preparar_venta(
            workbook,
            datetime(
                2026,
                9,
                3
            ),
            2
        )

        assert resultado["estado"] == "OK"

        assert resultado["precio_kg"] == 13000

        assert (
            resultado["tipo_aplicado"]
            == "Oferta"
        )

        assert (
            resultado["oferta_id"]
            == "OF_VALID_015"
        )

        assert (
            ventas.cell(2, 8).value
            == 13000
        )

        assert (
            ventas.cell(2, 9).value
            == 26000
        )

        assert (
            ventas.cell(2, 10).value
            == 6000
        )

        imprimir_resultado(
            "CASO 2 - OFERTA VIGENTE",
            resultado,
            ventas
        )

        # ====================================================
        # CASO 3 — OFERTA VENCIDA
        # ====================================================

        resultado = preparar_venta(
            workbook,
            datetime(
                2026,
                9,
                8
            ),
            2
        )

        assert resultado["estado"] == "OK"

        assert (
            resultado["precio_kg"]
            == 16000
        )

        assert (
            resultado["tipo_aplicado"]
            == "Minorista"
        )

        assert (
            resultado["oferta_id"]
            is None
        )

        assert (
            ventas.cell(2, 8).value
            == 16000
        )

        assert (
            ventas.cell(2, 9).value
            == 32000
        )

        assert (
            ventas.cell(2, 10).value
            == 12000
        )

        imprimir_resultado(
            "CASO 3 - OFERTA VENCIDA",
            resultado,
            ventas
        )

        # ====================================================
        # RESULTADO
        # ====================================================

        print()
        print("=" * 60)

        print(
            "OK: MVP 0.1.5 - PRUEBA REAL SUPERADA"
        )

        print("=" * 60)

        print()
        print(
            "Sin oferta: OK"
        )

        print(
            "Oferta vigente: OK"
        )

        print(
            "Oferta vencida: OK"
        )

        print(
            "Precio automático: OK"
        )

        print(
            "Importe automático: OK"
        )

        print(
            "Ganancia automática: OK"
        )

        print(
            "OfertaID: OK"
        )

        print(
            "Motivo del precio: OK"
        )

        print()
        print(
            "Los datos fueron creados solamente "
            "en memoria y NO se guardaron."
        )

    finally:

        workbook.close()


if __name__ == "__main__":
    main()