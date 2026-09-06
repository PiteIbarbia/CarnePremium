from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# RUTAS
# ============================================================

SCRIPT_DIR = Path(__file__).resolve().parent

SCRIPTS_DIR = (
    SCRIPT_DIR.parent
    / "Scripts"
)

sys.path.insert(
    0,
    str(SCRIPTS_DIR)
)


from config import MASTER_FILE

from ventas import calcular_venta


# ============================================================
# DATOS DE PRUEBA
# ============================================================

TEST_CLIENT = {
    "id": "CLI_TEST_001",
    "nombre": "Cliente Prueba",
    "telefono": "1111111111",
    "direccion": "Direccion de Prueba 123",
    "tipo_cliente": "Minorista",
    "estado": "Activo",
}


TEST_PRODUCT = {
    "id": "PRO_TEST_001",
    "nombre": "Asado Prueba",
    "categoria": "Vacuno",
    "costo": 10000,
    "precio_minorista": 16000,
    "precio_mayorista": None,
    "precio_especial": None,
    "estado": "Activo",
}


# ============================================================
# PREPARAR DATOS TEMPORALES
# ============================================================

def preparar_datos_prueba(workbook):

    clientes = workbook[
        "03_CLIENTES"
    ]

    productos = workbook[
        "04_PRODUCTOS"
    ]

    # --------------------------------------------------------
    # Cliente temporal
    # --------------------------------------------------------

    cliente_encontrado = False

    for fila in range(
        2,
        clientes.max_row + 1
    ):

        nombre = clientes.cell(
            fila,
            2
        ).value

        if (
            str(nombre).strip().lower()
            == TEST_CLIENT["nombre"].lower()
        ):

            cliente_encontrado = True
            break

    if not cliente_encontrado:

        # Buscar primera fila vacía
        fila = 2

        while (
            fila <= clientes.max_row
            and clientes.cell(fila, 2).value
        ):

            fila += 1

        clientes.cell(
            fila,
            1
        ).value = TEST_CLIENT["id"]

        clientes.cell(
            fila,
            2
        ).value = TEST_CLIENT["nombre"]

        clientes.cell(
            fila,
            3
        ).value = TEST_CLIENT["telefono"]

        clientes.cell(
            fila,
            4
        ).value = TEST_CLIENT["direccion"]

        clientes.cell(
            fila,
            5
        ).value = TEST_CLIENT["tipo_cliente"]

        clientes.cell(
            fila,
            6
        ).value = TEST_CLIENT["estado"]


    # --------------------------------------------------------
    # Producto temporal
    # --------------------------------------------------------

    producto_encontrado = False

    for fila in range(
        2,
        productos.max_row + 1
    ):

        nombre = productos.cell(
            fila,
            2
        ).value

        if (
            str(nombre).strip().lower()
            == TEST_PRODUCT["nombre"].lower()
        ):

            producto_encontrado = True
            break

    if not producto_encontrado:

        fila = 2

        while (
            fila <= productos.max_row
            and productos.cell(fila, 2).value
        ):

            fila += 1

        productos.cell(
            fila,
            1
        ).value = TEST_PRODUCT["id"]

        productos.cell(
            fila,
            2
        ).value = TEST_PRODUCT["nombre"]

        productos.cell(
            fila,
            3
        ).value = TEST_PRODUCT["categoria"]

        productos.cell(
            fila,
            4
        ).value = TEST_PRODUCT["costo"]

        productos.cell(
            fila,
            5
        ).value = TEST_PRODUCT["precio_minorista"]

        productos.cell(
            fila,
            6
        ).value = TEST_PRODUCT["precio_mayorista"]

        productos.cell(
            fila,
            7
        ).value = TEST_PRODUCT["precio_especial"]

        productos.cell(
            fila,
            8
        ).value = TEST_PRODUCT["estado"]


# ============================================================
# VALIDAR RESULTADO
# ============================================================

def validar_resultado(resultado):

    if resultado.get(
        "estado"
    ) != "OK":

        raise SystemExit(
            f"ERROR: estado = "
            f"{resultado.get('estado')}"
        )


    # --------------------------------------------------------
    # Valores esperados
    # --------------------------------------------------------

    esperado_precio = 16000

    esperado_importe = 32000

    esperado_ganancia = 12000

    esperado_porcentaje = 0.375


    # --------------------------------------------------------
    # Precio
    # --------------------------------------------------------

    if resultado["precio_kg"] != esperado_precio:

        raise SystemExit(
            f"ERROR: precio esperado "
            f"{esperado_precio}, "
            f"obtenido "
            f"{resultado['precio_kg']}"
        )


    # --------------------------------------------------------
    # Importe
    # --------------------------------------------------------

    if resultado["importe"] != esperado_importe:

        raise SystemExit(
            f"ERROR: importe esperado "
            f"{esperado_importe}, "
            f"obtenido "
            f"{resultado['importe']}"
        )


    # --------------------------------------------------------
    # Ganancia
    # --------------------------------------------------------

    if resultado["ganancia"] != esperado_ganancia:

        raise SystemExit(
            f"ERROR: ganancia esperada "
            f"{esperado_ganancia}, "
            f"obtenida "
            f"{resultado['ganancia']}"
        )


    # --------------------------------------------------------
    # Ganancia %
    # --------------------------------------------------------

    porcentaje = (
        resultado[
            "ganancia_porcentaje"
        ]
    )

    if abs(
        porcentaje
        - esperado_porcentaje
    ) > 0.000001:

        raise SystemExit(
            "ERROR: porcentaje de "
            "ganancia incorrecto. "
            f"Esperado: "
            f"{esperado_porcentaje}; "
            f"obtenido: "
            f"{porcentaje}"
        )


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 60)
    print("PRUEBA AUTOMATICA - MOTOR DE VENTAS")
    print("=" * 60)
    print()

    # --------------------------------------------------------
    # Verificar Excel
    # --------------------------------------------------------

    if not MASTER_FILE.exists():

        raise SystemExit(
            "ERROR: no existe el Excel maestro:\n"
            f"{MASTER_FILE}"
        )


    # --------------------------------------------------------
    # Abrir Excel
    # --------------------------------------------------------

    workbook = load_workbook(
        MASTER_FILE,
        data_only=False
    )


    try:

        print(
            "[1/3] Creando datos de prueba "
            "en memoria..."
        )

        preparar_datos_prueba(
            workbook
        )


        print(
            "[2/3] Procesando venta..."
        )

        resultado = calcular_venta(

            workbook,

            nombre_cliente=(
                "Cliente Prueba"
            ),

            nombre_producto=(
                "Asado Prueba"
            ),

            tipo_precio=(
                "Minorista"
            ),

            kg=2
        )


        print()
        print(
            "RESULTADO DE LA VENTA"
        )
        print("-" * 40)


        for clave, valor in (
            resultado.items()
        ):

            print(
                f"{clave}: {valor}"
            )


        print()
        print(
            "[3/3] Validando resultado..."
        )


        validar_resultado(
            resultado
        )


        print()
        print("=" * 60)
        print(
            "OK: PRUEBA AUTOMATICA SUPERADA"
        )
        print("=" * 60)
        print()

        print(
            "Cliente: Cliente Prueba"
        )

        print(
            "Producto: Asado Prueba"
        )

        print(
            "Kg: 2"
        )

        print(
            "Precio/Kg: $16.000"
        )

        print(
            "Importe: $32.000"
        )

        print(
            "Costo: $20.000"
        )

        print(
            "Ganancia: $12.000"
        )

        print(
            "Ganancia %: 37,5%"
        )

        print()
        print(
            "Los datos de prueba NO se "
            "guardaron en el Excel."
        )


    finally:

        workbook.close()


if __name__ == "__main__":

    main()