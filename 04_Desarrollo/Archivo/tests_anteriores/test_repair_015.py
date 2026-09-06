from __future__ import annotations

from pathlib import Path
import hashlib
import json
import sys

from openpyxl import load_workbook


# ============================================================
# RUTAS
# ============================================================

TEST_DIR = (
    Path(__file__)
    .resolve()
    .parent
)

SCRIPTS_DIR = (
    TEST_DIR.parent
    / "Scripts"
)

sys.path.insert(
    0,
    str(SCRIPTS_DIR)
)

from config import MASTER_FILE


STATE_FILE = (
    MASTER_FILE.parent
    / ".build_state.json"
)

EXPECTED_VERSION = "0.1.5"


# ============================================================
# HASH
# ============================================================

def sha256_file(
    path: Path
):

    sha = hashlib.sha256()

    with path.open(
        "rb"
    ) as file:

        for chunk in iter(
            lambda: file.read(
                1024 * 1024
            ),
            b""
        ):

            sha.update(
                chunk
            )

    return sha.hexdigest()


# ============================================================
# ERROR
# ============================================================

def fail(
    message
):

    print()
    print("=" * 60)
    print("TEST DE REPARACION FALLIDO")
    print("=" * 60)
    print(message)
    print()

    raise SystemExit(1)


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 60)
    print(
        "CARNE PREMIUM ERP - TEST REPAIR 0.1.5"
    )
    print("=" * 60)
    print()

    # ========================================================
    # ARCHIVO
    # ========================================================

    if not MASTER_FILE.exists():

        fail(
            f"No existe:\n{MASTER_FILE}"
        )

    # ========================================================
    # ESTADO
    # ========================================================

    if not STATE_FILE.exists():

        fail(
            "No existe .build_state.json."
        )

    try:

        state = json.loads(
            STATE_FILE.read_text(
                encoding="utf-8"
            )
        )

    except Exception as exc:

        fail(
            f"Estado inválido:\n{exc}"
        )

    # ========================================================
    # VERSION
    # ========================================================

    if (
        state.get("version")
        != EXPECTED_VERSION
    ):

        fail(
            "La versión del estado "
            "no es 0.1.5."
        )

    # ========================================================
    # HASH
    # ========================================================

    actual_hash = sha256_file(
        MASTER_FILE
    )

    if (
        actual_hash
        != state.get(
            "sha256"
        )
    ):

        fail(
            "El SHA-256 del maestro "
            "no coincide con el estado."
        )

    print(
        "[1/5] Maestro e integridad: OK"
    )

    # ========================================================
    # EXCEL
    # ========================================================

    workbook = load_workbook(
        MASTER_FILE,
        data_only=False
    )

    try:

        # ====================================================
        # VERSION
        # ====================================================

        version = workbook[
            "12_VERSION"
        ]

        if (
            version["B2"].value
            != EXPECTED_VERSION
        ):

            fail(
                "12_VERSION!B2 "
                "no es 0.1.5."
            )

        print(
            "[2/5] Versión: OK"
        )

        # ====================================================
        # PRODUCTOS
        # ====================================================

        products = workbook[
            "04_PRODUCTOS"
        ]

        expected_products = [

            "ProductoID",
            "Producto",
            "Categoria",
            "Costo",
            "PrecioMinorista",
            "PrecioMayorista",
            "PrecioEspecial",
            "Estado",
            "Observaciones",
        ]

        actual_products = [

            products.cell(
                1,
                column
            ).value

            for column
            in range(
                1,
                10
            )
        ]

        if (
            actual_products
            != expected_products
        ):

            fail(
                "04_PRODUCTOS todavía "
                "no tiene A:I correctamente."
            )

        # J debe estar vacío
        if products["J1"].value not in (
            None,
            ""
        ):

            fail(
                "04_PRODUCTOS!J1 "
                "todavía contiene datos."
            )

        for row in range(
            2,
            502
        ):

            if products.cell(
                row,
                10
            ).value not in (
                None,
                ""
            ):

                fail(
                    f"04_PRODUCTOS!J{row} "
                    "contiene datos."
                )

        table = products.tables.get(
            "ProductosTable"
        )

        if (
            table
            and
            table.ref
            != "A1:I501"
        ):

            fail(
                "ProductosTable no es "
                "A1:I501."
            )

        print(
            "[3/5] 04_PRODUCTOS reparado: OK"
        )

        # ====================================================
        # DASHBOARD
        # ====================================================

        dashboard = workbook[
            "10_DASHBOARD"
        ]

        expected = {

            "B3":
                '=COUNTIF(\'05_VENTAS\'!$A$2:$A$501,"<>")',

            "B4":
                '=SUM(\'05_VENTAS\'!$I$2:$I$501)',

            "B5":
                '=SUM(\'05_VENTAS\'!$J$2:$J$501)',

            "B6":
                '=IFERROR(B5/B4,"")',

            "B7":
                '=SUM(\'05_VENTAS\'!$G$2:$G$501)',

            "B8":
                '=IFERROR(B4/B3,"")',

            "B9":
                '=COUNTIF(\'05_VENTAS\'!$Q$2:$Q$501,"<>OK")-COUNTIF(\'05_VENTAS\'!$Q$2:$Q$501,"")',
        }

        for address, formula in (
            expected.items()
        ):

            if (
                dashboard[address].value
                != formula
            ):

                fail(
                    f"Dashboard!{address} "
                    "tiene una fórmula incorrecta."
                )

        print(
            "[4/5] Dashboard reparado: OK"
        )

        # ====================================================
        # OFERTAS
        # ====================================================

        offers = workbook[
            "23_OFERTAS"
        ]

        expected_offers = [

            "OfertaID",
            "ProductoID",
            "Producto",
            "NombreOferta",
            "PrecioOferta",
            "Descuento%",
            "FechaInicio",
            "FechaFin",
            "MinKg",
            "Estado",
            "Prioridad",
        ]

        actual_offers = [

            offers.cell(
                1,
                column
            ).value

            for column
            in range(
                1,
                12
            )
        ]

        if (
            actual_offers
            != expected_offers
        ):

            fail(
                "23_OFERTAS tiene "
                "estructura incorrecta."
            )

        print(
            "[5/5] 23_OFERTAS: OK"
        )

        print()
        print("=" * 60)
        print(
            "OK: TEST REPAIR 0.1.5 SUPERADO"
        )
        print("=" * 60)
        print()
        print(
            "La inconsistencia de Productos "
            "está corregida."
        )
        print(
            "Las fórmulas del Dashboard "
            "están corregidas."
        )
        print(
            "23_OFERTAS mantiene su estructura."
        )
        print(
            "El maestro coincide con "
            "su SHA-256."
        )
        print()

    finally:

        workbook.close()


if __name__ == "__main__":
    main()