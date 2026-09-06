from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import json
import shutil
import tempfile
import os

from openpyxl import load_workbook


# ============================================================
# CONFIGURACIÓN
# ============================================================

PROJECT_ROOT = (
    Path(__file__)
    .resolve()
    .parents[2]
)

MASTER_FILE = (
    PROJECT_ROOT
    / "01_ERP"
    / "CarnePremiumERP.xlsx"
)

BACKUP_DIR = (
    PROJECT_ROOT
    / "02_Backups"
)

STATE_FILE = (
    MASTER_FILE.parent
    / ".build_state.json"
)

VERSION = "0.1.5"

BUILD_NAME = (
    "MVP 0.1.5 - Ofertas automáticas en Ventas"
)

MAX_ROWS = 501


# ============================================================
# HASH
# ============================================================

def sha256_file(
    path: Path
) -> str:

    digest = hashlib.sha256()

    with path.open(
        "rb"
    ) as file:

        for chunk in iter(
            lambda: file.read(
                1024 * 1024
            ),
            b""
        ):

            digest.update(
                chunk
            )

    return digest.hexdigest()


# ============================================================
# BACKUP
# ============================================================

def create_backup():

    if not MASTER_FILE.exists():
        return None

    BACKUP_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    timestamp = (
        datetime.now()
        .strftime(
            "%Y%m%d_%H%M%S"
        )
    )

    backup = (
        BACKUP_DIR
        / (
            "CarnePremiumERP_"
            "before_repair_0.1.5_"
            f"{timestamp}.xlsx"
        )
    )

    shutil.copy2(
        MASTER_FILE,
        backup
    )

    return backup


# ============================================================
# VALIDAR CABECERAS
# ============================================================

def get_headers(
    ws,
    end_column
):

    return [
        ws.cell(
            row=1,
            column=column
        ).value
        for column
        in range(
            1,
            end_column + 1
        )
    ]


# ============================================================
# REPARAR 04_PRODUCTOS
# ============================================================

def repair_products(
    workbook
):

    ws = workbook[
        "04_PRODUCTOS"
    ]

    expected_headers = [

        "ProductoID",
        "Producto",
        "Categoria",
        "Costo",
        "PrecioMinorista",
        "PrecioMayorista",
        "PrecioEspecial",
        "Estado",
        "Observaciones",
    ]

    current_headers = get_headers(
        ws,
        10
    )

    # --------------------------------------------------------
    # Caso normal esperado actual:
    #
    # A:I = estructura correcta
    # J   = Observaciones duplicado
    # --------------------------------------------------------

    if (
        current_headers[0:9]
        != expected_headers
    ):

        raise RuntimeError(
            "04_PRODUCTOS no tiene la "
            "estructura A:I esperada.\n"
            f"Encontrado: {current_headers[0:9]}"
        )

    # --------------------------------------------------------
    # La columna J tiene que estar vacía.
    #
    # Nunca eliminamos datos desconocidos
    # automáticamente.
    # --------------------------------------------------------

    duplicated_header = (
        ws.cell(
            1,
            10
        ).value
    )

    if (
        duplicated_header
        not in (
            None,
            "",
            "Observaciones"
        )
    ):

        raise RuntimeError(
            "04_PRODUCTOS!J1 contiene "
            "un valor inesperado:\n"
            f"{duplicated_header}"
        )

    non_empty_j = []

    for row in range(
        2,
        MAX_ROWS + 1
    ):

        value = ws.cell(
            row,
            10
        ).value

        if value not in (
            None,
            ""
        ):

            non_empty_j.append(
                (row, value)
            )

    if non_empty_j:

        raise RuntimeError(
            "NO SE REPARA 04_PRODUCTOS "
            "AUTOMÁTICAMENTE porque la "
            "columna J contiene datos.\n"
            f"Primeros datos: "
            f"{non_empty_j[:5]}"
        )

    # --------------------------------------------------------
    # Limpiar J.
    # --------------------------------------------------------

    for row in range(
        1,
        MAX_ROWS + 1
    ):

        ws.cell(
            row,
            10
        ).value = None

    # --------------------------------------------------------
    # Tabla Productos
    # --------------------------------------------------------

    for table in ws.tables.values():

        if table.name == "ProductosTable":

            table.ref = (
                "A1:I501"
            )

    return True


# ============================================================
# REPARAR DASHBOARD
# ============================================================

def repair_dashboard(
    workbook
):

    ws = workbook[
        "10_DASHBOARD"
    ]

    # --------------------------------------------------------
    # A3:A9
    # --------------------------------------------------------

    expected_labels = [
        "Ventas registradas",
        "Facturación",
        "Ganancia",
        "Margen",
        "Kg vendidos",
        "Ticket promedio",
        "Ventas con error",
    ]

    for index, label in enumerate(
        expected_labels,
        start=3
    ):

        ws.cell(
            index,
            1
        ).value = label

    # --------------------------------------------------------
    # Fórmulas correctas
    # --------------------------------------------------------

    ws["B3"] = (
        '=COUNTIF('
        '\'05_VENTAS\'!'
        '$A$2:$A$501,'
        '"<>"'
        ')'
    )

    ws["B4"] = (
        '=SUM('
        '\'05_VENTAS\'!'
        '$I$2:$I$501'
        ')'
    )

    ws["B5"] = (
        '=SUM('
        '\'05_VENTAS\'!'
        '$J$2:$J$501'
        ')'
    )

    # Margen = Ganancia / Facturación
    ws["B6"] = (
        '=IFERROR('
        'B5/B4,'
        '"")'
    )

    ws["B7"] = (
        '=SUM('
        '\'05_VENTAS\'!'
        '$G$2:$G$501'
        ')'
    )

    # Ticket promedio = Facturación / cantidad de ventas
    ws["B8"] = (
        '=IFERROR('
        'B4/B3,'
        '"")'
    )

    ws["B9"] = (
        '=COUNTIF('
        '\'05_VENTAS\'!'
        '$Q$2:$Q$501,'
        '"<>OK"'
        ')'
        '-COUNTIF('
        '\'05_VENTAS\'!'
        '$Q$2:$Q$501,'
        '""'
        ')'
    )

    # --------------------------------------------------------
    # Formatos
    # --------------------------------------------------------

    ws["B4"].number_format = (
        '$ #,##0.00'
    )

    ws["B5"].number_format = (
        '$ #,##0.00'
    )

    ws["B6"].number_format = (
        '0.0%'
    )

    ws["B7"].number_format = (
        '0.000'
    )

    ws["B8"].number_format = (
        '$ #,##0.00'
    )

    return True


# ============================================================
# REPARAR 23_OFERTAS
# ============================================================

def repair_offers(
    workbook
):

    ws = workbook[
        "23_OFERTAS"
    ]

    expected_headers = [

        "OfertaID",
        "ProductoID",
        "Producto",
        "NombreOferta",
        "PrecioOferta",
        "Descuento%",
        "FechaInicio",
        "FechaFin",
        "MinKg",
        "Estado",
        "Prioridad",
    ]

    current = get_headers(
        ws,
        11
    )

    if current != expected_headers:

        raise RuntimeError(
            "23_OFERTAS no tiene la "
            "estructura esperada.\n"
            f"Encontrado: {current}"
        )

    # Fechas por día
    for row in range(
        2,
        MAX_ROWS + 1
    ):

        ws.cell(
            row,
            7
        ).number_format = (
            "dd/mm/yyyy"
        )

        ws.cell(
            row,
            8
        ).number_format = (
            "dd/mm/yyyy"
        )

        ws.cell(
            row,
            5
        ).number_format = (
            '$ #,##0.00'
        )

        ws.cell(
            row,
            6
        ).number_format = (
            '0.0%'
        )

        ws.cell(
            row,
            9
        ).number_format = (
            '0.000'
        )

    for table in ws.tables.values():

        if table.name == "OfertasTable":

            table.ref = (
                "A1:K501"
            )

    return True


# ============================================================
# VALIDACIÓN ESTRUCTURAL FINAL
# ============================================================

def validate_repair(
    workbook
):

    # --------------------------------------------------------
    # 04_PRODUCTOS
    # --------------------------------------------------------

    products = workbook[
        "04_PRODUCTOS"
    ]

    product_headers = get_headers(
        products,
        10
    )

    expected_products = [

        "ProductoID",
        "Producto",
        "Categoria",
        "Costo",
        "PrecioMinorista",
        "PrecioMayorista",
        "PrecioEspecial",
        "Estado",
        "Observaciones",
    ]

    if (
        product_headers[:9]
        != expected_products
    ):

        raise RuntimeError(
            "04_PRODUCTOS quedó con "
            "cabeceras incorrectas."
        )

    if product_headers[9] not in (
        None,
        ""
    ):

        raise RuntimeError(
            "04_PRODUCTOS todavía "
            "tiene una cabecera en J1."
        )

    for row in range(
        2,
        MAX_ROWS + 1
    ):

        if products.cell(
            row,
            10
        ).value not in (
            None,
            ""
        ):

            raise RuntimeError(
                "04_PRODUCTOS tiene datos "
                f"inesperados en J{row}."
            )

    products_table = (
        products.tables.get(
            "ProductosTable"
        )
    )

    if (
        products_table
        and
        products_table.ref
        != "A1:I501"
    ):

        raise RuntimeError(
            "ProductosTable debe ocupar "
            "A1:I501."
        )

    # --------------------------------------------------------
    # DASHBOARD
    # --------------------------------------------------------

    dashboard = workbook[
        "10_DASHBOARD"
    ]

    expected_dashboard = {

        "B3":
            '=COUNTIF(\'05_VENTAS\'!$A$2:$A$501,"<>")',

        "B4":
            '=SUM(\'05_VENTAS\'!$I$2:$I$501)',

        "B5":
            '=SUM(\'05_VENTAS\'!$J$2:$J$501)',

        "B6":
            '=IFERROR(B5/B4,"")',

        "B7":
            '=SUM(\'05_VENTAS\'!$G$2:$G$501)',

        "B8":
            '=IFERROR(B4/B3,"")',

        "B9":
            '=COUNTIF(\'05_VENTAS\'!$Q$2:$Q$501,"<>OK")-COUNTIF(\'05_VENTAS\'!$Q$2:$Q$501,"")',
    }

    for address, expected in (
        expected_dashboard.items()
    ):

        actual = dashboard[
            address
        ].value

        if actual != expected:

            raise RuntimeError(
                f"Dashboard!{address} "
                "no coincide.\n"
                f"Esperado: {expected}\n"
                f"Actual: {actual}"
            )

    # --------------------------------------------------------
    # 23_OFERTAS
    # --------------------------------------------------------

    offers = workbook[
        "23_OFERTAS"
    ]

    expected_offers = [

        "OfertaID",
        "ProductoID",
        "Producto",
        "NombreOferta",
        "PrecioOferta",
        "Descuento%",
        "FechaInicio",
        "FechaFin",
        "MinKg",
        "Estado",
        "Prioridad",
    ]

    if get_headers(
        offers,
        11
    ) != expected_offers:

        raise RuntimeError(
            "23_OFERTAS tiene "
            "cabeceras incorrectas."
        )


# ============================================================
# ESCRITURA ATÓMICA
# ============================================================

def save_atomically(
    workbook
):

    MASTER_FILE.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    fd, temp_name = (
        tempfile.mkstemp(
            prefix="CarnePremiumERP_repair_",
            suffix=".xlsx",
            dir=str(
                MASTER_FILE.parent
            )
        )
    )

    os.close(fd)

    temp_path = Path(
        temp_name
    )

    try:

        workbook.save(
            temp_path
        )

        # Reabrir para comprobar que
        # el archivo realmente quedó válido.
        check = load_workbook(
            temp_path,
            data_only=False
        )

        try:

            validate_repair(
                check
            )

        finally:

            check.close()

        backup = (
            create_backup()
        )

        if backup:

            print(
                f"Backup creado:\n{backup}"
            )

        temp_path.replace(
            MASTER_FILE
        )

        STATE_FILE.write_text(

            json.dumps(

                {
                    "version":
                        VERSION,

                    "build":
                        BUILD_NAME,

                    "repair":
                        "0.1.5",

                    "master_file":
                        str(
                            MASTER_FILE
                        ),

                    "sha256":
                        sha256_file(
                            MASTER_FILE
                        ),

                    "repaired_at":
                        datetime.now().isoformat(
                            timespec="seconds"
                        ),
                },

                indent=2,

                ensure_ascii=False
            ),

            encoding="utf-8"
        )

    finally:

        if temp_path.exists():

            try:

                temp_path.unlink()

            except PermissionError:

                pass


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 60)
    print(
        "CARNE PREMIUM ERP - REPARACION 0.1.5"
    )
    print("=" * 60)
    print()

    if not MASTER_FILE.exists():

        raise SystemExit(
            f"No existe:\n{MASTER_FILE}"
        )

    print(
        "[1/5] Abriendo maestro..."
    )

    workbook = load_workbook(
        MASTER_FILE,
        data_only=False
    )

    try:

        print(
            "[2/5] Reparando 04_PRODUCTOS..."
        )

        repair_products(
            workbook
        )

        print(
            "[3/5] Reparando 10_DASHBOARD..."
        )

        repair_dashboard(
            workbook
        )

        print(
            "[4/5] Verificando 23_OFERTAS..."
        )

        repair_offers(
            workbook
        )

        print(
            "[5/5] Validando y guardando..."
        )

        validate_repair(
            workbook
        )

        save_atomically(
            workbook
        )

    finally:

        workbook.close()

    print()
    print("=" * 60)
    print(
        "OK: REPARACION 0.1.5 COMPLETADA"
    )
    print("=" * 60)
    print()
    print(
        "04_PRODUCTOS: OK"
    )
    print(
        "Dashboard: OK"
    )
    print(
        "23_OFERTAS: OK"
    )
    print(
        "Backup: OK"
    )
    print(
        "SHA-256: OK"
    )
    print(
        f"Maestro: {MASTER_FILE}"
    )
    print()


if __name__ == "__main__":
    main()