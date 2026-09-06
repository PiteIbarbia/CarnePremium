from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

from config import MASTER_FILE, SHEETS

HEADER = {
    "fill": "#7A3E13",
    "font": {"bold": True, "color": "#FFFFFF"},
    "alignment": "center",
}


def ensure_workbook(path: Path):
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def set_headers(ws, headers):
    for col, header in enumerate(headers, 1):
        ws.cell(1, col).value = header
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7A3E13")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def setup_base(wb):
    for name in SHEETS:
        if name not in wb.sheetnames:
            wb.create_sheet(name)

    # Parameters
    p = wb["01_PARAMETROS"]
    p["A1"], p["B1"] = "Parametro", "Valor"
    p["A2"], p["B2"] = "Empresa", "Carne Premium"
    p["A3"], p["B3"] = "Moneda", "ARS"
    p["A4"], p["B4"] = "TipoPrecioDefault", "Minorista"
    p["A5"], p["B5"] = "MedioPagoDefault", "Efectivo"
    p["A6"], p["B6"] = "IVA", 21
    set_headers(p, ["Parametro", "Valor"])
    p.column_dimensions["A"].width = 24
    p.column_dimensions["B"].width = 24

    # Catalogs
    c = wb["02_CATALOGOS"]
    c["A1"], c["B1"], c["C1"] = "TiposPrecio", "TiposCliente", "MediosPago"
    for r, value in enumerate(["Minorista", "Mayorista", "Especial"], start=2):
        c.cell(r, 1).value = value
    for r, value in enumerate(["Minorista", "Mayorista", "Especial"], start=2):
        c.cell(r, 2).value = value
    for r, value in enumerate(["Efectivo", "Transferencia", "Cuenta Corriente", "Tarjeta", "Otro"], start=2):
        c.cell(r, 3).value = value
    set_headers(c, ["TiposPrecio", "TiposCliente", "MediosPago"])

    # Clients
    cl = wb["03_CLIENTES"]
    set_headers(cl, ["ClienteID", "Nombre", "Telefono", "Direccion", "TipoCliente", "Estado", "Observaciones"])
    cl["A2"] = '=IF(B2="","","CLI"&TEXT(ROW()-1,"000000"))'
    for r in range(3, 502):
        cl.cell(r, 1).value = f'=IF(B{r}="","","CLI"&TEXT(ROW()-1,"000000"))'
    for r in range(2, 502):
        cl.cell(r, 6).value = "Activo"
    cl.data_validations.add(DataValidation(type="list", formula1="'02_CATALOGOS'!$B$2:$B$4"), "E2:E501")
    cl.data_validations.add(DataValidation(type="list", values=["Activo", "Inactivo"]), "F2:F501")
    widths = [14, 28, 16, 28, 16, 14, 28]
    for i, width in enumerate(widths, 1):
        cl.column_dimensions[chr(64+i)].width = width
    cl.freeze_panes.freeze_rows(1)

    # Products
    pr = wb["04_PRODUCTOS"]
    set_headers(pr, ["ProductoID", "Producto", "Categoria", "Costo", "PrecioMinorista", "PrecioMayorista", "PrecioEspecial", "Estado", "Observaciones"])
    for r in range(2, 502):
        pr.cell(r, 1).value = f'=IF(B{r}="","","PRO"&TEXT(ROW()-1,"000000"))'
        pr.cell(r, 8).value = "Activo"
    pr.data_validations.add(DataValidation(type="list", values=["Activo", "Inactivo"]), "H2:H501")
    pr.column_dimensions["A"].width = 14
    pr.column_dimensions["B"].width = 28
    pr.column_dimensions["C"].width = 18
    for col in ["D", "E", "F", "G"]:
        pr.column_dimensions[col].width = 18
    pr.column_dimensions["H"].width = 14
    pr.column_dimensions["I"].width = 30
    pr.freeze_panes.freeze_rows(1)

    # Sales
    ve = wb["05_VENTAS"]
    set_headers(ve, ["VentaID", "Fecha", "Hora", "ClienteID", "TipoPrecio", "ProductoID", "Kg", "PrecioKg", "Importe", "Ganancia$", "Ganancia%", "MedioPago", "Vendedor", "Observaciones"])
    for r in range(2, 502):
        ve.cell(r, 1).value = f'=IF(OR(D{r}="",F{r}=""),"","VEN"&TEXT(ROW()-1,"000000"))'
        ve.cell(r, 8).value = f'=IFERROR(IF(E{r}="Mayorista",IF(VLOOKUP(F{r},\'04_PRODUCTOS\'!$A$2:$G$501,6,FALSE)="",VLOOKUP(F{r},\'04_PRODUCTOS\'!$A$2:$G$501,5,FALSE),VLOOKUP(F{r},\'04_PRODUCTOS\'!$A$2:$G$501,6,FALSE)),IF(E{r}="Especial",IF(VLOOKUP(F{r},\'04_PRODUCTOS\'!$A$2:$G$501,7,FALSE)="",VLOOKUP(F{r},\'04_PRODUCTOS\'!$A$2:$G$501,5,FALSE),VLOOKUP(F{r},\'04_PRODUCTOS\'!$A$2:$G$501,7,FALSE)),VLOOKUP(F{r},\'04_PRODUCTOS\'!$A$2:$G$501,5,FALSE))),"")'
        ve.cell(r, 9).value = f'=IF(AND(G{r}<>"",H{r}<>""),G{r}*H{r},"")'
        ve.cell(r, 10).value = f'=IFERROR(I{r}-(G{r}*VLOOKUP(F{r},\'04_PRODUCTOS\'!$A$2:$G$501,4,FALSE)),"")'
        ve.cell(r, 11).value = f'=IFERROR(J{r}/I{r},"")'
    ve.data_validations.add(DataValidation(type="list", formula1="'03_CLIENTES'!$A$2:$A$501"), "D2:D501")
    ve.data_validations.add(DataValidation(type="list", formula1="'02_CATALOGOS'!$A$2:$A$4"), "E2:E501")
    ve.data_validations.add(DataValidation(type="list", formula1="'04_PRODUCTOS'!$A$2:$A$501"), "F2:F501")
    ve.data_validations.add(DataValidation(type="list", formula1="'02_CATALOGOS'!$C$2:$C$6"), "L2:L501")
    for col in ["G"]:
        for r in range(2, 502):
            ve[f"{col}{r}"].number_format = "0.000"
    for col in ["H", "I", "J"]:
        for r in range(2, 502):
            ve[f"{col}{r}"].number_format = '$ #,##0.00'
    for r in range(2, 502):
        ve[f"K{r}"].number_format = '0.0%'
    ve.freeze_panes.freeze_rows(1)

    # Inicio / version
    inicio = wb["00_INICIO"]
    inicio["A1"] = "CARNE PREMIUM ERP"
    inicio["A2"] = "MVP 0.1.1"
    inicio["A4"] = "Prueba"
    inicio["B4"] = "Clientes -> Productos -> Ventas"
    inicio["A5"] = "Compatibilidad"
    inicio["B5"] = "Excel 2019 + Google Sheets"

    v = wb["12_VERSION"]
    v["A1"], v["B1"] = "Versión", "0.1.1"
    v["A2"], v["B2"] = "Build", "MVP Clientes + Productos + Ventas"
    v["A3"], v["B3"] = "Estado", "Prueba funcional"


def main():
    MASTER_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb = ensure_workbook(MASTER_FILE)
    setup_base(wb)
    wb.save(MASTER_FILE)
    print(f"ERP generado: {MASTER_FILE}")


if __name__ == "__main__":
    main()
