import sys
from pathlib import Path

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "Scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

from config import MASTER_FILE, SHEETS
from openpyxl import load_workbook

wb = load_workbook(MASTER_FILE, data_only=False)
missing = [s for s in SHEETS if s not in wb.sheetnames]
if missing:
    raise SystemExit(f"Faltan hojas: {missing}")

required = {
    "03_CLIENTES": ["ClienteID", "Nombre", "TipoCliente", "Estado"],
    "04_PRODUCTOS": ["ProductoID", "Producto", "Costo", "PrecioMinorista", "PrecioMayorista", "PrecioEspecial"],
    "05_VENTAS": ["VentaID", "ClienteID", "TipoPrecio", "ProductoID", "Kg", "PrecioKg", "Importe", "Ganancia$", "Ganancia%"],
}
for sheet, headers in required.items():
    current = [c.value for c in wb[sheet][1]]
    missing_headers = [h for h in headers if h not in current]
    if missing_headers:
        raise SystemExit(f"{sheet}: faltan encabezados {missing_headers}")

print("OK: MVP 0.1.1 validado: Clientes + Productos + Ventas.")
print(f"Archivo: {MASTER_FILE}")
