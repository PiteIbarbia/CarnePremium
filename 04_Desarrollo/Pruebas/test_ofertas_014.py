from __future__ import annotations

from datetime import datetime

import sys
from pathlib import Path

from openpyxl import Workbook


TEST_DIR = (
    Path(__file__).resolve().parent
)

SCRIPTS_DIR = (
    TEST_DIR.parent
    / "Scripts"
)

sys.path.insert(
    0,
    str(SCRIPTS_DIR)
)

from motor_comercial import (
    determinar_precio_venta
)


# ============================================================
# CREAR WORKBOOK DE PRUEBA
# ============================================================

def crear_workbook():

    workbook = Workbook()

    # --------------------------------------------------------
    # Eliminar hoja inicial
    # --------------------------------------------------------

    workbook.remove(
        workbook.active
    )

    # --------------------------------------------------------
    # CLIENTES
    # --------------------------------------------------------

    clientes = (
        workbook.create_sheet(
            "03_CLIENTES"
        )
    )

    clientes.append([
        "ClienteID",
        "Nombre",
        "Telefono",
        "Direccion",
        "TipoCliente",
        "Estado",
        "Observaciones"
    ])

    clientes.append([
        "CLI_TEST_001",
        "Cliente Prueba",
        "1111111111",
        "Direccion",
        "Minorista",
        "Activo",
        ""
    ])

    # --------------------------------------------------------
    # PRODUCTOS
    # --------------------------------------------------------

    productos = (
        workbook.create_sheet(
            "04_PRODUCTOS"
        )
    )

    productos.append([
        "ProductoID",
        "Producto",
        "Categoria",
        "Costo",
        "PrecioMinorista",
        "PrecioMayorista",
        "PrecioEspecial",
        "Estado",
        "Observaciones"
    ])

    productos.append([
        "PRO_TEST_001",
        "Asado Prueba",
        "Vacuno",
        10000,
        16000,
        14500,
        None,
        "Activo",
        ""
    ])

    # --------------------------------------------------------
    # OFERTAS
    # --------------------------------------------------------

    ofertas = (
        workbook.create_sheet(
            "23_OFERTAS"
        )
    )

    ofertas.append([
        "OfertaID",
        "ProductoID",
        "NombreOferta",
        "PrecioOferta",
        "Descuento%",
        "FechaInicio",
        "FechaFin",
        "MinKg",
        "Estado",
        "Prioridad"
    ])

    # Oferta 1
    ofertas.append([
        "OF001",
        "PRO_TEST_001",
        "Oferta Agosto",
        13500,
        None,
        datetime(2026, 8, 1),
        datetime(2026, 8, 31),
        1,
        "Activo",
        10
    ])

    # Oferta 2 - prioridad mayor
    ofertas.append([
        "OF002",
        "PRO_TEST_001",
        "Oferta Premium",
        13000,
        None,
        datetime(2026, 8, 15),
        datetime(2026, 8, 31),
        2,
        "Activo",
        20
    ])

    return workbook


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 60)
    print(
        "MVP 0.1.4 - TEST OFERTAS"
    )
    print("=" * 60)
    print()

    workbook = crear_workbook()

    try:

        # ====================================================
        # CASO 1
        # ====================================================

        print(
            "CASO 1 - FECHA VIGENTE / KG 1"
        )

        resultado = determinar_precio_venta(

            workbook,

            "Cliente Prueba",

            "Asado Prueba",

            datetime(
                2026,
                8,
                10
            ),

            1
        )

        print(
            resultado
        )

        assert (
            resultado["precio_kg"]
            == 13500
        )


        # ====================================================
        # CASO 2
        # ====================================================

        print()
        print(
            "CASO 2 - PRIORIDAD MAYOR"
        )

        resultado = determinar_precio_venta(

            workbook,

            "Cliente Prueba",

            "Asado Prueba",

            datetime(
                2026,
                8,
                20
            ),

            2
        )

        print(
            resultado
        )

        assert (
            resultado["precio_kg"]
            == 13000
        )

        assert (
            resultado["oferta_id"]
            == "OF002"
        )


        # ====================================================
        # CASO 3
        # ====================================================

        print()
        print(
            "CASO 3 - ULTIMO DIA INCLUSIVO"
        )

        resultado = determinar_precio_venta(

            workbook,

            "Cliente Prueba",

            "Asado Prueba",

            datetime(
                2026,
                8,
                31
            ),

            2
        )

        print(
            resultado
        )

        assert (
            resultado["precio_kg"]
            == 13000
        )


        # ====================================================
        # CASO 4
        # ====================================================

        print()
        print(
            "CASO 4 - DIA POSTERIOR"
        )

        resultado = determinar_precio_venta(

            workbook,

            "Cliente Prueba",

            "Asado Prueba",

            datetime(
                2026,
                9,
                1
            ),

            2
        )

        print(
            resultado
        )

        assert (
            resultado["precio_kg"]
            == 16000
        )


        # ====================================================
        # CASO 5
        # ====================================================

        print()
        print(
            "CASO 5 - KG MINIMO NO CUMPLIDO"
        )

        resultado = determinar_precio_venta(

            workbook,

            "Cliente Prueba",

            "Asado Prueba",

            datetime(
                2026,
                8,
                20
            ),

            1
        )

        print(
            resultado
        )

        # OF002 requiere 2 kg.
        # OF001 queda disponible porque requiere 1 kg.
        assert (
            resultado["precio_kg"]
            == 13500
        )


        # ====================================================
        # RESULTADO
        # ====================================================

        print()
        print("=" * 60)
        print(
            "OK: MVP 0.1.4 SUPERADO"
        )
        print("=" * 60)
        print()

        print(
            "Oferta vigente: OK"
        )

        print(
            "Fecha final inclusiva: OK"
        )

        print(
            "Prioridad de ofertas: OK"
        )

        print(
            "Kg mínimo: OK"
        )

        print(
            "Fallback a precio normal: OK"
        )

        print()

    finally:

        workbook.close()


if __name__ == "__main__":

    main()