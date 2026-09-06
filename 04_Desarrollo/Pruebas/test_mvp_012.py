from __future__ import annotations

from pathlib import Path
import sys

from openpyxl import load_workbook


TEST_DIR = Path(__file__).resolve().parent
SCRIPTS_DIR = TEST_DIR.parent / "Scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

from config import MASTER_FILE
from ventas_excel import procesar_venta_excel
from dashboard import preparar_dashboard


def preparar_cliente(workbook):
    ws = workbook["03_CLIENTES"]

    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 2).value or "").strip().lower() == "cliente prueba":
            return

    row = 2
    while row <= ws.max_row and ws.cell(row, 2).value:
        row += 1

    ws.cell(row, 1).value = "CLI_TEST_001"
    ws.cell(row, 2).value = "Cliente Prueba"
    ws.cell(row, 5).value = "Minorista"
    ws.cell(row, 6).value = "Activo"


def preparar_producto(workbook):
    ws = workbook["04_PRODUCTOS"]

    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, 2).value or "").strip().lower() == "asado prueba":
            return

    row = 2
    while row <= ws.max_row and ws.cell(row, 2).value:
        row += 1

    ws.cell(row, 1).value = "PRO_TEST_001"
    ws.cell(row, 2).value = "Asado Prueba"
    ws.cell(row, 3).value = "Vacuno"
    ws.cell(row, 4).value = 10000
    ws.cell(row, 5).value = 16000
    ws.cell(row, 6).value = None
    ws.cell(row, 7).value = None
    ws.cell(row, 8).value = "Activo"


def preparar_venta(workbook):
    ws = workbook["05_VENTAS"]
    row = 2

    ws.cell(row, 4).value = "Cliente Prueba"
    ws.cell(row, 5).value = "Minorista"
    ws.cell(row, 6).value = "Asado Prueba"
    ws.cell(row, 7).value = 2

    return row


def main():
    print("=" * 60)
    print("MVP 0.1.2 — PRUEBA DE INTEGRACIÓN")
    print("=" * 60)

    if not MASTER_FILE.exists():
        raise SystemExit(f"No existe: {MASTER_FILE}")

    workbook = load_workbook(MASTER_FILE, data_only=False)

    try:
        print("[1/4] Preparando datos de prueba...")
        preparar_cliente(workbook)
        preparar_producto(workbook)

        print("[2/4] Procesando venta...")
        fila = preparar_venta(workbook)
        resultado = procesar_venta_excel(workbook, fila)

        if resultado.get("estado") != "OK":
            raise SystemExit(
                f"ERROR: el motor devolvió {resultado.get('estado')}"
            )

        print("[3/4] Validando valores...")
        assert resultado["precio_kg"] == 16000
        assert resultado["importe"] == 32000
        assert resultado["ganancia"] == 12000
        assert abs(resultado["ganancia_porcentaje"] - 0.375) < 0.000001

        print("[4/4] Preparando Dashboard...")
        preparar_dashboard(workbook)

        ws = workbook["05_VENTAS"]

        assert ws.cell(fila, 15).value == "CLI_TEST_001"
        assert ws.cell(fila, 16).value == "PRO_TEST_001"
        assert ws.cell(fila, 17).value == "OK"

        print()
        print("OK: MVP 0.1.2 SUPERADO")
        print("Cliente por nombre: OK")
        print("Producto por nombre: OK")
        print("Precio automático: OK")
        print("Importe: OK")
        print("Ganancia: OK")
        print("IDs internos: OK")
        print("Estado de venta: OK")
        print("Dashboard: OK")
        print()
        print("La prueba funciona en memoria y NO guarda datos de prueba.")

    finally:
        workbook.close()


if __name__ == "__main__":
    main()
