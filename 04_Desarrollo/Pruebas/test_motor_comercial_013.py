from __future__ import annotations

from datetime import datetime

import sys
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# RUTAS
# ============================================================

TEST_DIR = (
    Path(__file__).resolve().parent
)

SCRIPTS_DIR = (
    TEST_DIR.parent
    / "Scripts"
)

sys.path.insert(
    0,
    str(SCRIPTS_DIR)
)


from config import MASTER_FILE

from motor_comercial import (
    determinar_precio_venta
)


# ============================================================
# PREPARAR CLIENTE
# ============================================================

def preparar_cliente(
    workbook
):
    hoja = workbook[
        "03_CLIENTES"
    ]

    nombre = "Cliente Prueba"

    for fila in range(
        2,
        hoja.max_row + 1
    ):

        if str(
            hoja.cell(
                fila,
                2
            ).value or ""
        ).strip().lower() == nombre.lower():

            return

    fila = 2

    while (
        fila <= hoja.max_row
        and hoja.cell(
            fila,
            2
        ).value
    ):

        fila += 1

    hoja.cell(
        fila,
        1
    ).value = "CLI_TEST_013"

    hoja.cell(
        fila,
        2
    ).value = "Cliente Prueba"

    hoja.cell(
        fila,
        3
    ).value = "1111111111"

    hoja.cell(
        fila,
        4
    ).value = "Direccion Prueba"

    hoja.cell(
        fila,
        5
    ).value = "Minorista"

    hoja.cell(
        fila,
        6
    ).value = "Activo"


# ============================================================
# PREPARAR PRODUCTO
# ============================================================

def preparar_producto(
    workbook
):
    hoja = workbook[
        "04_PRODUCTOS"
    ]

    nombre = "Asado Prueba"

    for fila in range(
        2,
        hoja.max_row + 1
    ):

        if str(
            hoja.cell(
                fila,
                2
            ).value or ""
        ).strip().lower() == nombre.lower():

            return

    fila = 2

    while (
        fila <= hoja.max_row
        and hoja.cell(
            fila,
            2
        ).value
    ):

        fila += 1

    hoja.cell(
        fila,
        1
    ).value = "PRO_TEST_013"

    hoja.cell(
        fila,
        2
    ).value = "Asado Prueba"

    hoja.cell(
        fila,
        3
    ).value = "Vacuno"

    hoja.cell(
        fila,
        4
    ).value = 10000

    hoja.cell(
        fila,
        5
    ).value = 16000

    hoja.cell(
        fila,
        6
    ).value = 14500

    hoja.cell(
        fila,
        7
    ).value = None

    hoja.cell(
        fila,
        8
    ).value = "Activo"


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 60)
    print(
        "MVP 0.1.3 - MOTOR COMERCIAL"
    )
    print("=" * 60)
    print()

    if not MASTER_FILE.exists():

        raise SystemExit(
            f"No existe:\n{MASTER_FILE}"
        )

    workbook = load_workbook(
        MASTER_FILE,
        data_only=False
    )

    try:

        # ----------------------------------------------------
        # Preparar datos temporales
        # ----------------------------------------------------

        preparar_cliente(
            workbook
        )

        preparar_producto(
            workbook
        )

        fecha = datetime(
            2026,
            8,
            31
        )

        # ----------------------------------------------------
        # CASO 1
        # ----------------------------------------------------

        print(
            "CASO 1 - MINORISTA"
        )

        resultado = (
            determinar_precio_venta(
                workbook,
                "Cliente Prueba",
                "Asado Prueba",
                fecha,
                2
            )
        )

        print(
            resultado
        )

        assert (
            resultado["estado"]
            == "OK"
        )

        assert (
            resultado["precio_kg"]
            == 16000
        )

        assert (
            resultado["tipo_aplicado"]
            == "Minorista"
        )

        # ----------------------------------------------------
        # CASO 2
        # ----------------------------------------------------

        print()
        print(
            "CASO 2 - MAYORISTA"
        )

        workbook[
            "03_CLIENTES"
        ].cell(
            2,
            5
        ).value = "Mayorista"

        resultado = (
            determinar_precio_venta(
                workbook,
                "Cliente Prueba",
                "Asado Prueba",
                fecha,
                2
            )
        )

        print(
            resultado
        )

        assert (
            resultado["precio_kg"]
            == 14500
        )

        assert (
            resultado["tipo_aplicado"]
            == "Mayorista"
        )

        # ----------------------------------------------------
        # CASO 3
        # ----------------------------------------------------

        print()
        print(
            "CASO 3 - MAYORISTA SIN PRECIO"
        )

        workbook[
            "04_PRODUCTOS"
        ].cell(
            2,
            6
        ).value = None

        resultado = (
            determinar_precio_venta(
                workbook,
                "Cliente Prueba",
                "Asado Prueba",
                fecha,
                2
            )
        )

        print(
            resultado
        )

        assert (
            resultado["precio_kg"]
            == 16000
        )

        # ----------------------------------------------------
        # CASO 4
        # ----------------------------------------------------

        print()
        print(
            "CASO 4 - OFERTA VIGENTE"
        )

        # No se crea todavía una oferta en Excel.
        # Esta parte se integrará contra 23_OFERTAS
        # una vez validada su estructura real.
        #
        # Por ahora esperamos fallback al precio base.

        resultado = (
            determinar_precio_venta(
                workbook,
                "Cliente Prueba",
                "Asado Prueba",
                fecha,
                2
            )
        )

        print(
            resultado
        )

        assert (
            resultado["precio_kg"]
            == 16000
        )

        # ----------------------------------------------------
        # FINAL
        # ----------------------------------------------------

        print()
        print("=" * 60)
        print(
            "OK: MOTOR COMERCIAL INTEGRADO BASE"
        )
        print("=" * 60)
        print()
        print(
            "Cliente por nombre: OK"
        )
        print(
            "Producto por nombre: OK"
        )
        print(
            "Minorista: OK"
        )
        print(
            "Mayorista: OK"
        )
        print(
            "Fallback a minorista: OK"
        )
        print(
            "No se modificó el Excel maestro."
        )
        print()

    finally:

        workbook.close()


if __name__ == "__main__":
    main()
    