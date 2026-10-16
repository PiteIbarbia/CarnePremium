from __future__ import annotations

import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

SCRIPT_DIR = Path(__file__).resolve().parents[1] / "Scripts"
sys.path.insert(0, str(SCRIPT_DIR))

from caja import preparar_caja, sincronizar_ventas_en_caja
from cuenta_corriente import sincronizar_cuenta_corriente
from motor_comercial import find_client, find_product
from ticket import preparar_ticket
from ventas_excel import procesar_todas_las_ventas
import builder
from builder import setup_sales, recover_legacy_sales_from_backups


def make_workbook():
    wb = Workbook()
    sales = wb.active
    sales.title = "05_VENTAS"
    sales.append([
        "VentaID", "Fecha", "Hora", "Cliente", "TipoPrecio", "Producto",
        "Kg", "PrecioKg", "Importe", "Ganancia$", "Ganancia%", "MedioPago",
        "Vendedor", "Observaciones", "_ClienteID", "_ProductoID", "EstadoVenta",
        "TipoAplicado", "MotivoPrecio", "OfertaID", "OfertaNombre",
        "PrecioEspecialID", "PagoRecibido", "_TicketLineaKey", "EstadoCobro",
        "SaldoTicket", "TipoRegistro", "FechaCobro",
    ])

    clients = wb.create_sheet("03_CLIENTES")
    clients.append(["ClienteID", "Nombre", "Telefono", "Direccion", "TipoCliente", "Estado"])
    clients.append(['=IF(B2="","","CLI"&TEXT(ROW()-1,"000000"))', "Cliente Demo", None, None, "Minorista", "Activo"])
    clients.append(['=IF(B3="","","CLI"&TEXT(ROW()-1,"000000"))', "Cliente Cuenta", None, None, "Minorista", "Activo"])

    # IDs de cliente/producto se generan por fórmulas en Excel. openpyxl lee
    # esas fórmulas como texto, y el motor debe resolver IDs estables por fila.
    products = wb.create_sheet("04_PRODUCTOS")
    products.append([
        "ProductoID", "Producto", "Categoria", "Costo", "PrecioMinorista",
        "PrecioMayorista", "PrecioEspecial", "Estado", "Observaciones",
    ])
    products.append([
        '=IF(B2="","","PRO"&TEXT(ROW()-1,"000000"))',
        "Asado", "Carne", 5000, 10000, 9000, 8500, "Activo", None,
    ])

    wb.create_sheet("06_TICKET")
    wb.create_sheet("07_CAJA")
    wb.create_sheet("09_CTACTE")

    rows = [
        ["VEN000001", datetime(2026, 10, 11, 10), datetime(2026, 10, 11, 10), "Cliente Demo", "Minorista", "Asado", 2, 10000, 20000, 3000, .15, "Efectivo", "Vendedor Demo", None, "CLI000001", "PRO000001", "OK", "Minorista", "Precio base", None, None, None, 15000, None, "SALDO PENDIENTE", 5000, "VENTA"],
        ["VEN000002", datetime(2026, 10, 11, 10, 5), datetime(2026, 10, 11, 10, 5), "Cliente Demo", "Minorista", "Vacío", 1.5, 10000, 15000, 2500, .1667, "Efectivo", "Vendedor Demo", None, "CLI000001", "PRO000002", "OK", "Minorista", "Precio base", None, None, None, 15000, None, "SALDO PENDIENTE", 5000, "VENTA"],
        ["VEN000003", datetime(2024, 4, 15, 10), datetime(2024, 4, 15, 10), "Cliente Demo", "Minorista", "Bife", .5, 10000, 5000, 1000, .2, "Efectivo", "Vendedor Demo", None, "CLI000001", "PRO000003", "OK", "Minorista", "Precio base", None, None, None, None, None, "PENDIENTE", 5000, "VENTA"],
        ["VEN000004", datetime(2026, 10, 11, 10, 15), datetime(2026, 10, 11, 10, 15), "Cliente Desconocido", "Minorista", "Colita", .5, None, None, None, None, "Efectivo", "Vendedor Demo", None, None, None, "CLIENTE NO ENCONTRADO", None, None, None, None, None, None, None, None, None, "VENTA"],
        ["VEN000005", datetime(2026, 10, 11, 10, 20), datetime(2026, 10, 11, 10, 20), "Cliente Cuenta", "Minorista", "Matambre", 1, 10000, 10000, 2000, .2, "Cuenta Corriente", "Vendedor Demo", None, "CLI000002", "PRO000004", "OK", "Minorista", "Precio base", None, None, None, None, None, "PENDIENTE", 10000, "VENTA"],
        ["", datetime(2026, 10, 12, 12), datetime(2026, 10, 12, 12), "Cliente Cuenta", None, None, None, None, None, None, None, "Transferencia", "Vendedor Demo", "Cobro parcial de saldo", "CLI000002", None, "COBRO OK", None, None, None, None, None, 3000, None, None, None, "COBRO"],
    ]
    for row in rows:
        sales.append(row)
    return wb



def test_order_payment_does_not_remove_sale_or_ticket():
    """Un pedido pendiente genera ticket; cobrarlo no lo transforma ni lo borra."""
    wb = make_workbook()
    try:
        sales = wb["05_VENTAS"]
        # Aislar la fila de prueba de los datos auxiliares del fixture.
        for row in range(3, sales.max_row + 1):
            for col in range(1, 29):
                sales.cell(row, col).value = None

        sales["AA2"] = "COBRO"  # Simula la clasificación errónea de una venta heredada.
        sales["L2"] = None
        sales["W2"] = None
        setup_sales(wb)
        procesar_todas_las_ventas(wb, 2, 2)

        assert sales["AA2"].value == "VENTA", sales["AA2"].value
        assert sales["Q2"].value == "OK", sales["Q2"].value
        assert sales["H2"].value == 10000
        assert sales["I2"].value == 20000

        preparar_ticket(wb)
        ticket = wb["06_TICKET"]
        ticket_id = "T-20261011-CLI000001"
        listed_before = {
            ticket.cell(row, 7).value
            for row in range(2, 502)
            if ticket.cell(row, 7).value not in (None, "")
        }
        assert ticket_id in listed_before, listed_before

        # Cobro posterior: la fecha de venta y el importe comercial se conservan.
        sales["L2"] = "Efectivo"
        sales["W2"] = 20000
        sales["AB2"] = datetime(2026, 10, 12, 13, 30)
        procesar_todas_las_ventas(wb, 2, 2)
        preparar_ticket(wb)
        cash_result = sincronizar_ventas_en_caja(wb)
        sincronizar_cuenta_corriente(wb)

        assert sales["AA2"].value == "VENTA"
        assert sales["H2"].value == 10000
        assert sales["I2"].value == 20000
        listed_after = {
            ticket.cell(row, 7).value
            for row in range(2, 502)
            if ticket.cell(row, 7).value not in (None, "")
        }
        assert ticket_id in listed_after, listed_after

        cash = wb["07_CAJA"]
        assert cash_result["ingresos_cobrados"] == 20000, cash_result
        assert cash["B2"].value == datetime(2026, 10, 12, 13, 30), cash["B2"].value
        assert cash["G2"].value == 20000

        account = wb["09_CTACTE"]
        rows_by_day = {
            account.cell(row, 1).value: row
            for row in range(2, account.max_row + 1)
            if account.cell(row, 3).value == "Cliente Demo"
        }
        day_sale = rows_by_day[datetime(2026, 10, 11).date()]
        day_paid = rows_by_day[datetime(2026, 10, 12).date()]
        assert account.cell(day_sale, 9).value == 20000
        assert account.cell(day_paid, 9).value == 0
        assert account.cell(day_paid, 10).value == "PAGADO / OK"
    finally:
        wb.close()



def test_backup_recovers_cleared_historical_sale_values():
    """Si H:I fueron borradas por un COBRO erróneo, usar backup o no guardar."""
    wb = make_workbook()
    original_master = builder.MASTER_FILE
    try:
        sales = wb["05_VENTAS"]
        sales["AA2"] = "COBRO"
        sales["Q2"] = "COBRO OK"
        for cell in ("H2", "I2", "J2", "K2"):
            sales[cell] = None

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            backup_dir = root / "02_Backups"
            backup_dir.mkdir(parents=True, exist_ok=True)
            builder.MASTER_FILE = root / "01_ERP" / "CarnePremiumERP.xlsx"

            backup = Workbook()
            old_sales = backup.active
            old_sales.title = "05_VENTAS"
            for col in range(1, sales.max_column + 1):
                old_sales.cell(1, col).value = sales.cell(1, col).value
                old_sales.cell(2, col).value = sales.cell(2, col).value
            old_sales["AA2"] = "VENTA"
            old_sales["Q2"] = "OK"
            old_sales["H2"] = 10000
            old_sales["I2"] = 20000
            old_sales["J2"] = 3000
            old_sales["K2"] = .15
            old_sales["O2"] = "CLI000001"
            old_sales["P2"] = "PRO000001"
            old_sales["R2"] = "Minorista"
            old_sales["S2"] = "Precio base"
            backup_path = backup_dir / "CarnePremiumERP_before_v0.5.1_test.xlsx"
            backup.save(backup_path)
            backup.close()

            result = recover_legacy_sales_from_backups(wb)
            assert result["restauradas"] == 1, result
            assert sales["AA2"].value == "VENTA"
            assert sales["Q2"].value == "OK"
            assert sales["H2"].value == 10000
            assert sales["I2"].value == 20000
    finally:
        builder.MASTER_FILE = original_master
        wb.close()




def test_legacy_auto_reversals_are_audited_not_fake_expenses():
    """Las anulaciones heredadas no deben crear egresos artificiales ni borrar el importe."""
    wb = make_workbook()
    try:
        cash = preparar_caja(wb)
        cash["A2"] = "MOV000099"
        cash["B2"] = datetime(2026, 10, 9, 9, 0)
        cash["C2"] = "ANULACION"
        cash["D2"] = "Anulación automática: VEN000999"
        cash["E2"] = "VEN000999"
        cash["F2"] = "Efectivo"
        cash["G2"] = 0
        cash["H2"] = 12500
        cash["I2"] = -12500
        cash["J2"] = "VENTA_AUTO_ANULADA"

        result = sincronizar_ventas_en_caja(wb)
        revisions = [
            row for row in range(2, cash.max_row + 1)
            if cash.cell(row, 10).value == "VENTA_AUTO_REVISION"
            and cash.cell(row, 5).value == "VEN000999"
        ]
        assert len(revisions) == 1, result
        row = revisions[0]
        cc = {str(cash.cell(1, col).value): col for col in range(1, cash.max_column + 1) if cash.cell(1, col).value}
        assert cash.cell(row, cc["Tipo"]).value == "AJUSTE PENDIENTE"
        assert cash.cell(row, cc["Ingreso"]).value == 0
        assert cash.cell(row, cc["Egreso"]).value == 0
        assert cash.cell(row, cc["Importe original"]).value == 12500
        assert "revisar" in cash.cell(row, cc["Observación / ajuste"]).value.lower()

        # En este fixture solo hay ingresos positivos; una reversión heredada no debe
        # producir un saldo acumulado negativo simulando una salida de efectivo.
        running = 0
        for movement_row in range(2, result["movimientos"] + 2):
            running += float(cash.cell(movement_row, cc["Ingreso"]).value or 0)
            running -= float(cash.cell(movement_row, cc["Egreso"]).value or 0)
            assert running >= 0, (movement_row, running)
        assert result["ajustes_pendientes"] == 1
    finally:
        wb.close()


def test_duplicate_auto_cash_movement_is_kept_for_audit():
    """Los duplicados se conservan marcados y no duplican el efectivo."""
    wb = make_workbook()
    try:
        cash = preparar_caja(wb)
        for row in (2, 3):
            cash.cell(row, 1).value = f"MOV{row:06d}"
            cash.cell(row, 2).value = datetime(2026, 10, 11, 10, 0)
            cash.cell(row, 3).value = "INGRESO"
            cash.cell(row, 4).value = "Cobro histórico duplicado"
            cash.cell(row, 5).value = "VEN000001"
            cash.cell(row, 6).value = "Efectivo"
            cash.cell(row, 7).value = 15000
            cash.cell(row, 8).value = 0
            cash.cell(row, 10).value = "VENTA_AUTO"
        # Una columna adicional de usuario debe moverse junto con su movimiento.
        cash["M2"] = "nota-movimiento-uno"
        cash["M3"] = "nota-movimiento-dos"

        result = sincronizar_ventas_en_caja(wb)
        cc = {str(cash.cell(1, col).value): col for col in range(1, cash.max_column + 1) if cash.cell(1, col).value}
        review_rows = [
            row for row in range(2, result["movimientos"] + 2)
            if cash.cell(row, cc["Origen"]).value == "VENTA_AUTO_REVISION"
            and "duplicado" in str(cash.cell(row, cc["Observación / ajuste"]).value).lower()
        ]
        active_rows = [
            row for row in range(2, result["movimientos"] + 2)
            if cash.cell(row, cc["Origen"]).value == "VENTA_AUTO"
            and cash.cell(row, cc["ReferenciaVenta"]).value == "VEN000001"
        ]
        income_total = sum(float(cash.cell(row, cc["Ingreso"]).value or 0) for row in range(2, result["movimientos"] + 2))
        expense_total = sum(float(cash.cell(row, cc["Egreso"]).value or 0) for row in range(2, result["movimientos"] + 2))
        assert len(review_rows) == 1, (review_rows, result)
        assert len(active_rows) == 1, active_rows
        assert income_total == 33000, (income_total, result)
        assert expense_total == 0, expense_total
        assert result["ajustes_pendientes"] == 1
        row_by_id = {
            str(cash.cell(row, 1).value): row
            for row in range(2, result["movimientos"] + 2)
        }
        assert cash.cell(row_by_id["MOV000002"], 13).value == "nota-movimiento-uno"
        assert cash.cell(row_by_id["MOV000003"], 13).value == "nota-movimiento-dos"
    finally:
        wb.close()

def main():
    test_backup_recovers_cleared_historical_sale_values()
    test_order_payment_does_not_remove_sale_or_ticket()
    test_legacy_auto_reversals_are_audited_not_fake_expenses()
    test_duplicate_auto_cash_movement_is_kept_for_audit()
    wb = make_workbook()
    try:
        # Regresión: nunca se deben copiar textos de fórmulas de ID a Ventas.
        assert find_client(wb, "Cliente Demo")["id"] == "CLI000001"
        assert find_client(wb, "Cliente Cuenta")["id"] == "CLI000002"
        assert find_product(wb, "Asado")["id"] == "PRO000001"

        first_cash = sincronizar_ventas_en_caja(wb)
        preparar_ticket(wb)
        first_account = sincronizar_cuenta_corriente(wb)

        assert first_cash["agregados"] == 3, first_cash
        assert first_cash["actualizados"] == 0, first_cash
        assert first_cash["ingresos_cobrados"] == 33000, first_cash

        cash = wb["07_CAJA"]
        by_ref = {
            str(cash.cell(row, 5).value): row
            for row in range(2, cash.max_row + 1)
            if cash.cell(row, 10).value == "VENTA_AUTO"
        }
        assert set(by_ref) == {"VEN000001", "VEN000002", "COB000006"}, by_ref
        cash_total = sum(float(cash.cell(row, 7).value or 0) for row in range(2, cash.max_row + 1))
        assert cash_total == 33000, cash_total

        ticket = wb["06_TICKET"]
        listed = [
            ticket.cell(row, 7).value
            for row in range(2, 502)
            if ticket.cell(row, 7).value not in (None, "")
        ]
        assert listed == [
            "T-20261011-CLI000001",
            "T-20240415-CLI000001",
            "T-20261011-CLI000002",
        ], listed
        assert ticket["B3"].value == "T-20261011-CLI000001"
        assert ticket["A9"].value.startswith("=IFERROR(INDEX(")
        assert ticket["B39"].value == "=SUM(B9:B38)"
        assert ticket["B39"].number_format == "0.000"
        assert ticket["D39"].value == "=SUM(D9:D38)"

        second_cash = sincronizar_ventas_en_caja(wb)
        preparar_ticket(wb)
        second_account = sincronizar_cuenta_corriente(wb)
        assert second_cash["agregados"] == 0, second_cash
        assert second_cash["actualizados"] == 3, second_cash
        assert second_cash["movimientos"] == 3, second_cash
        assert cash["N1"].value == "RESUMEN CAJA"
        assert cash["O2"].value == "=SUM(G2:G10000)"
        assert cash["O3"].value == "=SUM(H2:H10000)"
        assert cash["O4"].value == "=SUM(G2:G10000)-SUM(H2:H10000)"
        assert cash["O5"].value == '=COUNTIF(C2:C10000,"AJUSTE PENDIENTE")'
        assert cash["N6"].value == "Saldo inicial de efectivo"
        assert cash["O6"].value == "No incluido"

        # Editar un cobro existente actualiza el movimiento correspondiente, no crea otro.
        wb["05_VENTAS"]["W2"] = 16000
        third_cash = sincronizar_ventas_en_caja(wb)
        sincronizar_cuenta_corriente(wb)
        assert third_cash["agregados"] == 0, third_cash
        assert third_cash["movimientos"] == 3, third_cash
        assert cash.cell(by_ref["VEN000001"], 7).value == 16000

        # Cuenta Corriente suma ventas, cobros de ventas y cobros posteriores.
        account = wb["09_CTACTE"]
        account_rows = {
            (account.cell(row, 3).value, account.cell(row, 1).value): row
            for row in range(2, account.max_row + 1)
        }
        demo_day1 = account_rows[("Cliente Demo", datetime(2026, 10, 11).date())]
        account_day1 = account_rows[("Cliente Cuenta", datetime(2026, 10, 11).date())]
        account_day2 = account_rows[("Cliente Cuenta", datetime(2026, 10, 12).date())]
        assert account.cell(demo_day1, 9).value == 9000  # incluye los $5.000 pendientes de la venta histórica
        assert account.cell(account_day1, 9).value == 10000
        assert account.cell(account_day2, 9).value == 7000
        assert account.cell(account_day2, 10).value == "DEBE"

        print("OK: Ticket automático cliente/fecha, kilos totales, cobros, Caja idempotente y Cuenta Corriente.")
    finally:
        wb.close()


if __name__ == "__main__":
    main()
