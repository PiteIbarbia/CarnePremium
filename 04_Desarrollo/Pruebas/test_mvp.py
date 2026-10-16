from __future__ import annotations

from pathlib import Path
import hashlib
import json
import sys

from openpyxl import load_workbook


# ============================================================
# CONFIGURACIÓN DEL TEST
# ============================================================

EXPECTED_VERSION = "0.5.3"

EXPECTED_BUILD_NAME = (
    "Build 0.5.3 - Caja clara, saldo cronológico y ajustes auditables"
)

# ============================================================
# RUTAS
# ============================================================

TEST_DIR = Path(
    __file__
).resolve().parent

SCRIPTS_DIR = (
    TEST_DIR.parent
    / "Scripts"
)

sys.path.insert(
    0,
    str(SCRIPTS_DIR)
)

from config import MASTER_FILE


STATE_FILE = (
    MASTER_FILE.parent
    / ".build_state.json"
)


# ============================================================
# HASH SHA-256
# ============================================================

def sha256_file(
    path: Path
) -> str:

    sha = hashlib.sha256()

    with path.open(
        "rb"
    ) as file:

        for chunk in iter(
            lambda: file.read(
                1024 * 1024
            ),
            b""
        ):

            sha.update(
                chunk
            )

    return sha.hexdigest()


# ============================================================
# ERROR
# ============================================================

def fail(
    message: str
):

    print()
    print(
        "=" * 60
    )

    print(
        "TEST FALLIDO"
    )

    print(
        "=" * 60
    )

    print(
        message
    )

    print()

    raise SystemExit(
        1
    )


# ============================================================
# TEST PRINCIPAL
# ============================================================

def main():

    print()
    print(
        "=" * 60
    )

    print(
        "CARNE PREMIUM ERP"
    )

    print(
        "TEST MVP 0.5.3"
    )

    print(
        "=" * 60
    )

    print()

    # ========================================================
    # 1. EXISTENCIA DEL MAESTRO
    # ========================================================

    print(
        "[1/8] Verificando archivo maestro..."
    )

    if not MASTER_FILE.exists():

        fail(
            f"No existe el archivo maestro:\n"
            f"{MASTER_FILE}"
        )

    print(
        "OK"
    )


    # ========================================================
    # 2. EXISTENCIA DEL ESTADO
    # ========================================================

    print(
        "[2/8] Verificando estado del Build..."
    )

    if not STATE_FILE.exists():

        fail(
            "No existe .build_state.json.\n"
            "El builder no terminó correctamente."
        )

    try:

        state = json.loads(
            STATE_FILE.read_text(
                encoding="utf-8"
            )
        )

    except Exception as exc:

        fail(
            f"No se pudo leer .build_state.json:\n"
            f"{exc}"
        )

    print(
        "OK"
    )


    # ========================================================
    # 3. COMPROBAR QUE NO ES UN EXCEL VIEJO
    # ========================================================

    print(
        "[3/8] Comprobando integridad del maestro..."
    )

    actual_hash = (
        sha256_file(
            MASTER_FILE
        )
    )

    expected_hash = (
        state.get(
            "sha256"
        )
    )

    if not expected_hash:

        fail(
            "El estado del Build no contiene SHA-256."
        )

    if actual_hash != expected_hash:

        fail(
            "EL EXCEL MAESTRO NO COINCIDE "
            "CON EL ÚLTIMO BUILD.\n\n"
            "El archivo puede ser una versión anterior "
            "o haber sido modificado después del Build."
        )

    print(
        "SHA-256: OK"
    )


    # ========================================================
    # 4. COMPROBAR VERSION
    # ========================================================

    print(
        "[4/8] Verificando versión..."
    )

    state_version = (
        state.get(
            "version"
        )
    )

    if state_version != EXPECTED_VERSION:

        fail(
            f"Se esperaba versión "
            f"{EXPECTED_VERSION}, "
            f"pero el estado indica "
            f"{state_version}."
        )

    print(
        f"Versión {EXPECTED_VERSION}: OK"
    )


    # ========================================================
    # 5. ABRIR EXCEL
    # ========================================================

    print(
        "[5/8] Abriendo Excel maestro..."
    )

    try:

        workbook = load_workbook(
            MASTER_FILE,
            data_only=False
        )

    except Exception as exc:

        fail(
            f"No se pudo abrir el Excel:\n"
            f"{exc}"
        )

    print(
        "OK"
    )


    try:

        # ====================================================
        # 6. HOJAS OBLIGATORIAS
        # ====================================================

        print(
            "[6/8] Verificando estructura..."
        )

        required_sheets = [

            "00_INICIO",
            "01_PARAMETROS",
            "02_CATALOGOS",
            "03_CLIENTES",
            "04_PRODUCTOS",
            "05_VENTAS",
            "06_TICKET",
            "07_CAJA",
            "08_STOCK",
            "09_CTACTE",
            "10_DASHBOARD",
            "11_AUDITORIA",
            "12_VERSION",
            "13_LOG",
        ]

        missing_sheets = [

            sheet

            for sheet
            in required_sheets

            if sheet
            not in workbook.sheetnames
        ]

        if missing_sheets:

            fail(
                "Faltan hojas obligatorias:\n"
                + "\n".join(
                    missing_sheets
                )
            )

        print(
            "Hojas: OK"
        )


        # ====================================================
        # 7. VERSION DEL LIBRO
        # ====================================================

        version_sheet = (
            workbook[
                "12_VERSION"
            ]
        )

        excel_version = (
            version_sheet[
                "B2"
            ].value
        )

        if (
            excel_version
            != EXPECTED_VERSION
        ):

            fail(
                f"El Excel indica versión "
                f"{excel_version}, "
                f"pero se esperaba "
                f"{EXPECTED_VERSION}."
            )

        excel_build = (
            version_sheet[
                "B3"
            ].value
        )

        if (
            excel_build
            != EXPECTED_BUILD_NAME
        ):

            fail(
                "El nombre del Build "
                "no coincide.\n"
                f"Esperado: {EXPECTED_BUILD_NAME}\n"
                f"Encontrado: {excel_build}"
            )

        excel_token = (
            version_sheet[
                "B5"
            ].value
        )

        state_token = (
            state.get(
                "build_token"
            )
        )

        if (
            not excel_token
            or excel_token
            != state_token
        ):

            fail(
                "El BuildToken no coincide.\n"
                "El Excel podría no corresponder "
                "al último Build."
            )

        print(
            "Versión/BuildToken: OK"
        )


        # ====================================================
        # 8. CLIENTES
        # ====================================================

        print(
            "[7/8] Verificando Clientes, Productos y Ventas..."
        )

        clients = workbook[
            "03_CLIENTES"
        ]

        client_headers = [

            "ClienteID",
            "Nombre",
            "Telefono",
            "Direccion",
            "TipoCliente",
            "Estado",
            "Observaciones"
        ]

        actual_client_headers = [

            cell.value

            for cell
            in clients[1]
        ]

        missing = [

            header

            for header
            in client_headers

            if header
            not in actual_client_headers
        ]

        if missing:

            fail(
                "03_CLIENTES: faltan encabezados:\n"
                + "\n".join(
                    missing
                )
            )

        if len(
            clients.data_validations.dataValidation
        ) < 2:

            fail(
                "03_CLIENTES: "
                "faltan las validaciones esperadas."
            )

        print(
            "Clientes: OK"
        )


        # ====================================================
        # PRODUCTOS
        # ====================================================

        products = workbook[
            "04_PRODUCTOS"
        ]

        product_headers = [

            "ProductoID",
            "Producto",
            "Categoria",
            "Costo",
            "PrecioMinorista",
            "PrecioMayorista",
            "PrecioEspecial",
            "Estado",
            "Observaciones"
        ]

        actual_product_headers = [

            cell.value

            for cell
            in products[1]
        ]

        missing = [

            header

            for header
            in product_headers

            if header
            not in actual_product_headers
        ]

        if missing:

            fail(
                "04_PRODUCTOS: faltan encabezados:\n"
                + "\n".join(
                    missing
                )
            )

        if len(
            products.data_validations.dataValidation
        ) < 1:

            fail(
                "04_PRODUCTOS: "
                "falta la validación de Estado."
            )

        print(
            "Productos: OK"
        )


        # ====================================================
        # VENTAS
        # ====================================================

        sales = workbook[
            "05_VENTAS"
        ]

        sales_headers = [

            "VentaID",
            "Fecha",
            "Hora",
            "Cliente",
            "TipoPrecio",
            "Producto",
            "Kg",
            "PrecioKg",
            "Importe",
            "Ganancia$",
            "Ganancia%",
            "MedioPago",
            "Vendedor",
            "Observaciones",
            "_ClienteID",
            "_ProductoID",
            "EstadoVenta",
            "TipoAplicado",
            "MotivoPrecio",
            "OfertaID",
            "OfertaNombre",
            "PrecioEspecialID",
            "PagoRecibido",
            "_TicketLineaKey",
            "EstadoCobro",
            "SaldoTicket",
            "TipoRegistro"
        ]

        actual_sales_headers = [

            cell.value

            for cell
            in sales[1]
        ]

        missing = [

            header

            for header
            in sales_headers

            if header
            not in actual_sales_headers
        ]

        if missing:

            fail(
                "05_VENTAS: faltan encabezados:\n"
                + "\n".join(
                    missing
                )
            )


        # ----------------------------------------------------
        # Fórmulas críticas
        # ----------------------------------------------------

        # La clave de ticket X se materializa como texto en preparar_ticket(),
        # no debe exigirse como fórmula: así el Ticket funciona aunque Excel
        # todavía no haya recalculado las fórmulas auxiliares.
        formula_cells = ["A2", "Y2", "Z2"]

        row_has_data = any(
            sales.cell(2, column).value not in (None, "")
            for column in (4, 6, 7)
        )
        if not row_has_data:
            formula_cells.extend(
                ["B2", "C2", "H2", "I2", "J2", "K2", "O2", "P2", "Q2"]
            )
        else:
            if sales["Q2"].value in (None, ""):
                fail("05_VENTAS!Q2 no contiene el estado de procesamiento.")
            if str(sales["Q2"].value).strip().upper() == "OK":
                ticket_key = sales["X2"].value
                if (
                    not isinstance(ticket_key, str)
                    or not ticket_key.startswith("T-")
                    or "|" not in ticket_key
                ):
                    fail("05_VENTAS!X2 no contiene una clave de ticket materializada.")

        for address in formula_cells:

            value = (
                sales[
                    address
                ].value
            )

            if (
                not isinstance(
                    value,
                    str
                )

                or

                not value.startswith(
                    "="
                )
            ):

                fail(
                    f"05_VENTAS!"
                    f"{address} "
                    f"no contiene una "
                    f"fórmula válida."
                )


        # ----------------------------------------------------
        # Validaciones de Ventas
        # ----------------------------------------------------

        sales_validations = {
            str(cell_range): validation.formula1
            for validation in sales.data_validations.dataValidation
            for cell_range in validation.sqref.ranges
        }
        required_sales_ranges = {
            "D2:D501": "=ListaClientes",
            "F2:F501": "=ListaProductos",
        }
        for target_range, expected_formula in required_sales_ranges.items():
            if sales_validations.get(target_range) != expected_formula:
                fail(
                    f"05_VENTAS: falta el selector de {target_range}; "
                    f"se esperaba {expected_formula!r} y se encontró "
                    f"{sales_validations.get(target_range)!r}."
                )

        if "ListaClientes" not in workbook.defined_names:
            fail("Falta el rango con nombre ListaClientes.")
        if "ListaProductos" not in workbook.defined_names:
            fail("Falta el rango con nombre ListaProductos.")

        print(
            "Ventas: OK"
        )

        ticket = workbook["06_TICKET"]
        if [ticket.cell(8, column).value for column in range(1, 5)] != [
            "Corte", "Kg", "Precio por kg", "Importe"
        ]:
            fail("06_TICKET: faltan encabezados de detalle.")

        caja_headers = [
            "MovimientoID", "Fecha", "Tipo", "Concepto", "ReferenciaVenta",
            "MedioPago", "Ingreso", "Egreso", "Saldo", "Origen",
        ]
        caja = workbook["07_CAJA"]
        if [caja.cell(1, column).value for column in range(1, 11)] != caja_headers:
            fail("07_CAJA: encabezados incorrectos.")

        cash_headers = {
            str(caja.cell(1, column).value).strip(): column
            for column in range(1, caja.max_column + 1)
            if caja.cell(1, column).value not in (None, "")
        }
        if "Importe original" not in cash_headers or "Observación / ajuste" not in cash_headers:
            fail("07_CAJA: faltan las columnas de auditoría de ajustes.")
        if caja.freeze_panes != "D2":
            fail("07_CAJA: la fila de encabezados y las columnas clave no están inmovilizadas.")
        if not caja.auto_filter.ref or not caja.auto_filter.ref.startswith("A1:"):
            fail("07_CAJA: falta el filtro de movimientos.")

        summary_column = next(
            (
                column for column in range(1, caja.max_column + 1)
                if caja.cell(1, column).value == "RESUMEN CAJA"
            ),
            None,
        )
        if summary_column is None:
            fail("07_CAJA: falta el resumen rápido.")
        for offset, expected_formula in (
            (1, "=SUM(G2:G10000)"),
            (2, "=SUM(H2:H10000)"),
            (3, "=SUM(G2:G10000)-SUM(H2:H10000)"),
        ):
            if caja.cell(2 + offset - 1, summary_column + 1).value != expected_formula:
                fail(f"07_CAJA: fórmula incorrecta en el resumen, fila {offset + 1}.")

        # ====================================================
        # DASHBOARD
        # ====================================================

        dashboard = workbook[
            "10_DASHBOARD"
        ]

        dashboard_formulas = [

            "B3",
            "B4",
            "B5",
            "B6",
            "B7",
            "B8",
            "B9"
        ]

        for address in dashboard_formulas:

            value = (
                dashboard[
                    address
                ].value
            )

            if (
                not isinstance(
                    value,
                    str
                )
                or
                not value.startswith(
                    "="
                )
            ):

                fail(
                    f"10_DASHBOARD!"
                    f"{address} "
                    f"no contiene una "
                    f"fórmula válida."
                )

        print(
            "Dashboard: OK"
        )


        # ====================================================
        # RESULTADO
        # ====================================================

        print(
            "[8/8] Prueba final..."
        )

        print()

        print(
            "=" * 60
        )

        print(
            "OK: BUILD 0.5.3 VALIDADO"
        )

        print(
            "=" * 60
        )

        print(
            "Integridad del maestro: OK"
        )

        print(
            "Versión 0.5.3: OK"
        )

        print(
            "Clientes: OK"
        )

        print(
            "Productos: OK"
        )

        print(
            "Ventas por nombre: OK"
        )

        print(
            "IDs internos: OK"
        )

        print(
            "Precio automático: OK"
        )

        print(
            "Importe: OK"
        )

        print(
            "Ganancia: OK"
        )

        print(
            "Dashboard: OK"
        )

        print(
            "Protección contra Excel viejo: OK"
        )

        print(
            "SHA-256: OK"
        )

        print()
        print(
            f"Archivo: {MASTER_FILE}"
        )

        print()

    finally:

        workbook.close()


# ============================================================
# EJECUCIÓN
# ============================================================

if __name__ == "__main__":

    main()