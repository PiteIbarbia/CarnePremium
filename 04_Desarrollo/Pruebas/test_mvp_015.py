from __future__ import annotations

from datetime import datetime

import sys
from pathlib import Path

from openpyxl import Workbook

TEST_DIR = (
    Path(__file__).resolve().parent
)

SCRIPTS_DIR = (
    TEST_DIR.parent
    / "Scripts"
)

sys.path.insert(
    0,
    str(SCRIPTS_DIR)
)

from ventas_excel import (
    procesar_venta_excel
)


def crear_workbook():

    workbook = Workbook()

    workbook.remove(
        workbook.active
    )

    # ========================================================
    # CLIENTES
    # ========================================================

    clientes = workbook.create_sheet(
        "03_CLIENTES"
    )

    clientes.append([
        "ClienteID",
        "Nombre",
        "Telefono",
        "Direccion",
        "TipoCliente",
        "Estado",
        "Observaciones"
    ])

    clientes.append([
        "CLI015",
        "Cliente Prueba",
        "1111111111",
        "Direccion",
        "Minorista",
        "Activo",
        ""
    ])

    # ========================================================
    # PRODUCTOS
    # ========================================================

    productos = workbook.create_sheet(
        "04_PRODUCTOS"
    )

    productos.append([
        "ProductoID",
        "Producto",
        "Categoria",
        "Costo",
        "PrecioMinorista",
        "PrecioMayorista",
        "PrecioEspecial",
        "Estado",
        "Observaciones"
    ])

    productos.append([
        "PRO015",
        "Asado Prueba",
        "Vacuno",
        10000,
        16000,
        14500,
        14000,
        "Activo",
        ""
    ])

    # ========================================================
    # OFERTAS
    # ========================================================

    ofertas = workbook.create_sheet(
        "23_OFERTAS"
    )

    ofertas.append([
        "OfertaID",
        "ProductoID",
        "NombreOferta",
        "PrecioOferta",
        "Descuento%",
        "FechaInicio",
        "FechaFin",
        "MinKg",
        "Estado",
        "Prioridad"
    ])

    ofertas.append([
        "OF015",
        "PRO015",
        "Oferta Prueba",
        13000,
        None,
        datetime(
            2026,
            8,
            1
        ),
        datetime(
            2026,
            8,
            31
        ),
        2,
        "Activo",
        100
    ])

    # ========================================================
    # VENTAS
    # ========================================================

    ventas = workbook.create_sheet(
        "05_VENTAS"
    )

    ventas.append([
        "VentaID",
        "Fecha",
        "Hora",
        "Cliente",
        "TipoPrecio",
        "Producto",
        "Kg",
        "PrecioKg",
        "Importe",
        "Ganancia$",
        "Ganancia%",
        "MedioPago",
        "Vendedor",
        "Observaciones",
        "_ClienteID",
        "_ProductoID",
        "EstadoVenta",
        "TipoAplicado",
        "MotivoPrecio",
        "OfertaID",
        "OfertaNombre"
    ])

    return workbook


def main():

    print()
    print(
        "=" * 60
    )

    print(
        "MVP 0.1.5 - INTEGRACION OFERTAS + VENTAS"
    )

    print(
        "=" * 60
    )

    workbook = crear_workbook()

    try:

        # ====================================================
        # PREPARAR VENTA
        # ====================================================

        ventas = workbook[
            "05_VENTAS"
        ]

        ventas.cell(
            2,
            2
        ).value = datetime(
            2026,
            8,
            20
        )

        ventas.cell(
            2,
            4
        ).value = (
            "Cliente Prueba"
        )

        ventas.cell(
            2,
            5
        ).value = (
            "Minorista"
        )

        ventas.cell(
            2,
            6
        ).value = (
            "Asado Prueba"
        )

        ventas.cell(
            2,
            7
        ).value = 2

        # ====================================================
        # PROCESAR
        # ====================================================

        print()
        print(
            "[1/3] Procesando venta..."
        )

        resultado = (
            procesar_venta_excel(
                workbook,
                2
            )
        )

        print(
            resultado
        )

        # ====================================================
        # VALIDAR
        # ====================================================

        print()
        print(
            "[2/3] Validando oferta..."
        )

        assert (
            resultado["estado"]
            == "OK"
        )

        assert (
            resultado["precio_kg"]
            == 13000
        )

        assert (
            resultado["tipo_aplicado"]
            == "Oferta"
        )

        assert (
            resultado["oferta_id"]
            == "OF015"
        )

        assert (
            resultado["oferta_nombre"]
            == "Oferta Prueba"
        )

        # ====================================================
        # VALIDAR EXCEL
        # ====================================================

        print()
        print(
            "[3/3] Validando datos escritos..."
        )

        assert (
            ventas.cell(
                2,
                8
            ).value
            == 13000
        )

        assert (
            ventas.cell(
                2,
                9
            ).value
            == 26000
        )

        assert (
            ventas.cell(
                2,
                10
            ).value
            == 6000
        )

        assert (
            ventas.cell(
                2,
                17
            ).value
            == "OK"
        )

        assert (
            ventas.cell(
                2,
                18
            ).value
            == "Oferta"
        )

        assert (
            ventas.cell(
                2,
                19
            ).value
            == "Oferta Prueba"
        )

        assert (
            ventas.cell(
                2,
                20
            ).value
            == "OF015"
        )

        assert (
            ventas.cell(
                2,
                21
            ).value
            == "Oferta Prueba"
        )

        print()
        print(
            "=" * 60
        )

        print(
            "OK: MVP 0.1.5 SUPERADO"
        )

        print(
            "=" * 60
        )

        print()
        print(
            "Oferta automática: OK"
        )

        print(
            "Precio final: OK"
        )

        print(
            "Importe: OK"
        )

        print(
            "Ganancia: OK"
        )

        print(
            "OfertaID: OK"
        )

        print(
            "Motivo del precio: OK"
        )

        print()

    finally:

        workbook.close()


if __name__ == "__main__":

    main()