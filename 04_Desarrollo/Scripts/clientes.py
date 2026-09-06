from openpyxl import load_workbook


def obtener_cliente_por_nombre(
    workbook,
    nombre_cliente
):

    hoja = workbook[
        "03_CLIENTES"
    ]

    if not nombre_cliente:
        return None

    nombre_buscado = (
        str(nombre_cliente)
        .strip()
        .lower()
    )

    for fila in range(
        2,
        hoja.max_row + 1
    ):

        nombre = hoja.cell(
            fila,
            2
        ).value

        if not nombre:
            continue

        if (
            str(nombre)
            .strip()
            .lower()
            == nombre_buscado
        ):

            return {
                "id": hoja.cell(
                    fila,
                    1
                ).value,

                "nombre": nombre,

                "telefono": hoja.cell(
                    fila,
                    3
                ).value,

                "direccion": hoja.cell(
                    fila,
                    4
                ).value,

                "tipo_cliente": hoja.cell(
                    fila,
                    5
                ).value,

                "estado": hoja.cell(
                    fila,
                    6
                ).value,
            }

    return None