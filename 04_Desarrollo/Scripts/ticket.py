from __future__ import annotations

from datetime import date, datetime
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

MAX_SALES_ROW = 501
FIRST_LINE_ROW = 9
TICKET_LINE_COUNT = 30
TOTAL_ROW = FIRST_LINE_ROW + TICKET_LINE_COUNT
LIST_COLUMN = 7


def _column_map(ws):
    result = {}
    for column in range(1, ws.max_column + 1):
        value = ws.cell(1, column).value
        if value not in (None, ""):
            result[str(value).strip()] = column
    return result


def _text(value):
    return str(value or "").strip()


def _as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        raw = value.strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(raw[:10], fmt).date()
            except ValueError:
                continue
    return None


def _automatic_ticket_id(client_id, sale_date):
    """Ticket único por cliente y día, sin identificador manual en Ventas."""
    return f"T-{sale_date:%Y%m%d}-{client_id}"


def _client_id_for_name(workbook, client_name):
    """Resuelve un ID de cliente aun cuando la celda de ID tenga fórmula Excel."""
    if "03_CLIENTES" not in workbook.sheetnames:
        return ""
    target = _text(client_name).casefold()
    clients = workbook["03_CLIENTES"]
    for row in range(2, min(MAX_SALES_ROW, clients.max_row) + 1):
        name = _text(clients.cell(row, 2).value)
        if not name or name.casefold() != target:
            continue
        raw_id = clients.cell(row, 1).value
        if raw_id in (None, "") or (isinstance(raw_id, str) and raw_id.startswith("=")):
            return f"CLI{row - 1:06d}"
        return str(raw_id).strip()
    return ""


def preparar_ticket(workbook):
    """Prepara 06_TICKET y un selector con tickets de ventas aprobadas."""
    if "05_VENTAS" not in workbook.sheetnames:
        raise RuntimeError("Falta la hoja 05_VENTAS.")

    sales = workbook["05_VENTAS"]
    columns = _column_map(sales)
    required = [
        "VentaID", "Fecha", "Cliente", "_ClienteID", "Producto", "Kg",
        "PrecioKg", "Importe", "MedioPago", "Vendedor",
        "EstadoVenta", "PagoRecibido", "_TicketLineaKey",
    ]
    missing = [name for name in required if name not in columns]
    if missing:
        raise RuntimeError("05_VENTAS: faltan columnas para el ticket: " + ", ".join(missing))

    ticket_ids = []
    seen = set()
    for row in range(2, min(MAX_SALES_ROW, sales.max_row) + 1):
        if _text(sales.cell(row, columns["EstadoVenta"]).value).upper() != "OK":
            continue
        if not _text(sales.cell(row, columns["Cliente"]).value):
            continue
        if not _text(sales.cell(row, columns["Producto"]).value):
            continue
        try:
            kg = float(sales.cell(row, columns["Kg"]).value)
        except (TypeError, ValueError):
            continue
        if kg <= 0:
            continue
        client_id = _client_id_for_name(
            workbook, sales.cell(row, columns["Cliente"]).value
        )
        sale_date = _as_date(sales.cell(row, columns["Fecha"]).value)
        if not client_id or sale_date is None:
            continue
        ticket_id = _automatic_ticket_id(client_id, sale_date)
        if ticket_id not in seen:
            seen.add(ticket_id)
            ticket_ids.append(ticket_id)

    ws = workbook["06_TICKET"] if "06_TICKET" in workbook.sheetnames else workbook.create_sheet("06_TICKET")
    selected_before = _text(ws["B3"].value)

    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    for row in range(1, 45):
        for column in range(1, 5):
            ws.cell(row, column).value = None
    for row in range(1, MAX_SALES_ROW + 1):
        ws.cell(row, LIST_COLUMN).value = None
    ws.data_validations.dataValidation = []

    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1")
    ws["A1"] = "CARNE PREMIUM | TICKET DE VENTA"
    ws["A1"].font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="7A3E13")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    for address, label in {
        "A3": "TicketID", "A4": "Fecha", "A5": "Cliente",
        "A6": "Vendedor", "A7": "Medio de pago",
    }.items():
        ws[address] = label
        ws[address].font = Font(bold=True, color="7A3E13")

    selected = selected_before if selected_before in ticket_ids else (ticket_ids[0] if ticket_ids else None)
    ws["B3"] = selected
    ws["B3"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["B3"].font = Font(bold=True, color="7A3E13")

    ws["G1"] = "TicketID_Lista"
    for idx, ticket_id in enumerate(ticket_ids, start=2):
        ws.cell(idx, LIST_COLUMN).value = ticket_id
    ws.column_dimensions["G"].hidden = True
    validation = DataValidation(
        type="list", formula1="$G$2:$G$" + str(MAX_SALES_ROW), allow_blank=True
    )
    validation.errorTitle = "Ticket no válido"
    validation.error = "Elegí un TicketID de la lista."
    validation.promptTitle = "Seleccionar ticket"
    validation.prompt = "Elegí el ticket para consultar o imprimir."
    ws.add_data_validation(validation)
    validation.add(ws["B3"])

    helper_letter = get_column_letter(columns["_TicketLineaKey"])
    helper_range = "'05_VENTAS'!$" + helper_letter + "$2:$" + helper_letter + "$" + str(MAX_SALES_ROW)

    def detail_formula(header, row):
        letter = get_column_letter(columns[header])
        data_range = "'05_VENTAS'!$" + letter + "$2:$" + letter + "$" + str(MAX_SALES_ROW)
        occurrence = '"|"&ROWS($A$' + str(FIRST_LINE_ROW) + ':A' + str(row) + ')'
        return '=IFERROR(INDEX(' + data_range + ',MATCH($B$3&' + occurrence + ',' + helper_range + ',0)),"")'

    def header_formula(header):
        letter = get_column_letter(columns[header])
        data_range = "'05_VENTAS'!$" + letter + "$2:$" + letter + "$" + str(MAX_SALES_ROW)
        return '=IFERROR(INDEX(' + data_range + ',MATCH($B$3&"|1",' + helper_range + ',0)),"")'

    for address, header in {"B4": "Fecha", "B5": "Cliente", "B6": "Vendedor", "B7": "MedioPago"}.items():
        ws[address] = header_formula(header)
    ws["B4"].number_format = "dd/mm/yyyy"

    for column, label in enumerate(["Corte", "Kg", "Precio por kg", "Importe"], start=1):
        cell = ws.cell(8, column)
        cell.value = label
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7A3E13")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in range(FIRST_LINE_ROW, FIRST_LINE_ROW + TICKET_LINE_COUNT):
        for column, header in enumerate(["Producto", "Kg", "PrecioKg", "Importe"], start=1):
            ws.cell(row, column).value = detail_formula(header, row)
        ws.cell(row, 2).number_format = "0.000"
        ws.cell(row, 3).number_format = '$ #,##0.00'
        ws.cell(row, 4).number_format = '$ #,##0.00'

    # Totales del ticket: kilos e importe.
    # La columna B es Kg; C conserva la etiqueta TOTAL y D suma los importes.
    ws.cell(TOTAL_ROW, 2).value = f"=SUM(B{FIRST_LINE_ROW}:B{FIRST_LINE_ROW + TICKET_LINE_COUNT - 1})"
    ws.cell(TOTAL_ROW, 2).font = Font(bold=True, color="FFFFFF")
    ws.cell(TOTAL_ROW, 2).fill = PatternFill("solid", fgColor="7A3E13")
    ws.cell(TOTAL_ROW, 2).number_format = "0.000"
    ws.cell(TOTAL_ROW, 2).alignment = Alignment(horizontal="right")

    ws.cell(TOTAL_ROW, 3).value = "TOTAL"
    ws.cell(TOTAL_ROW, 3).font = Font(bold=True, color="FFFFFF")
    ws.cell(TOTAL_ROW, 3).fill = PatternFill("solid", fgColor="7A3E13")
    ws.cell(TOTAL_ROW, 4).value = f"=SUM(D{FIRST_LINE_ROW}:D{FIRST_LINE_ROW + TICKET_LINE_COUNT - 1})"
    ws.cell(TOTAL_ROW, 4).font = Font(bold=True, color="FFFFFF")
    ws.cell(TOTAL_ROW, 4).fill = PatternFill("solid", fgColor="7A3E13")
    ws.cell(TOTAL_ROW, 4).number_format = '$ #,##0.00'

    ws.merge_cells("A42:D44")
    ws["A42"] = (
        "El ticket se agrupa automáticamente por cliente y fecha. En W (PagoRecibido), "
        "registrá el importe efectivamente cobrado; podés distribuir pagos en una o más "
        "líneas del mismo cliente y día. EstadoCobro y SaldoTicket se calculan en Ventas. "
        "Ejecutá el ERP para refrescar Ticket, Caja y Cuenta Corriente."
    )
    ws["A42"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["A42"].font = Font(italic=True, color="666666", size=9)
    for col, width in {"A": 30, "B": 15, "C": 20, "D": 20}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A9"
    ws.print_area = f"A1:D{TOTAL_ROW}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.orientation = "portrait"
