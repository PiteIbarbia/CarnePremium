from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_MASTER = PROJECT_ROOT / "01_ERP" / "CarnePremiumERP.xlsx"
BACKUP_DIR = PROJECT_ROOT / "02_Backups"
SCRIPTS_DIR = PROJECT_ROOT / "04_Desarrollo" / "Scripts"
TESTS_DIR = PROJECT_ROOT / "04_Desarrollo" / "Pruebas"
STATE_NAME = ".build_state.json"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def valid_xlsx(path: Path) -> bool:
    try:
        with ZipFile(path, "r") as archive:
            return archive.testzip() is None
    except (OSError, BadZipFile):
        return False


def atomic_copy(source: Path, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{target.stem}_",
        suffix=target.suffix or ".tmp",
        dir=str(target.parent),
    )
    os.close(fd)
    temporary = Path(temp_name)
    try:
        shutil.copy2(source, temporary)
        if target.suffix.lower() == ".xlsx" and not valid_xlsx(temporary):
            raise RuntimeError("La copia no supera la validación XLSX.")
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)


def run_step(label: str, script: Path, env: dict[str, str]) -> bool:
    print()
    print("=" * 72)
    print(label)
    print("=" * 72)
    result = subprocess.run(
        [sys.executable, str(script)],
        cwd=PROJECT_ROOT,
        env=env,
        check=False,
    )
    if result.returncode != 0:
        print(f"ERROR: {script.name} terminó con código {result.returncode}.")
        return False
    return True



def report_pending_adjustments(path: Path) -> int:
    """Imprime, sin modificar el libro, las filas de Caja que requieren conciliación."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if "07_CAJA" not in workbook.sheetnames:
            print("AVISO: el libro generado no contiene 07_CAJA.")
            return 0

        ws = workbook["07_CAJA"]
        columns = {
            str(ws.cell(1, column).value).strip(): column
            for column in range(1, ws.max_column + 1)
            if ws.cell(1, column).value not in (None, "")
        }
        required = ("Tipo", "Fecha", "Concepto", "ReferenciaVenta", "Origen",
                    "Ingreso", "Egreso", "Importe original", "Observación / ajuste")
        missing = [name for name in required if name not in columns]
        if missing:
            print(
                "AVISO: no se pudieron detallar los ajustes; faltan columnas: "
                + ", ".join(missing)
            )
            return 0

        pending = []
        for row in range(2, ws.max_row + 1):
            kind = str(ws.cell(row, columns["Tipo"]).value or "").strip().upper()
            if kind == "AJUSTE PENDIENTE":
                pending.append(row)

        print()
        print("=" * 72)
        print(f"REVISIÓN DE CAJA — AJUSTES PENDIENTES: {len(pending)}")
        print("Estos movimientos quedan fuera del saldo hasta que se concilien.")
        print("=" * 72)
        if not pending:
            print("No hay ajustes pendientes en el libro generado.")
            return 0

        for index, row in enumerate(pending, start=1):
            def value(header: str):
                return ws.cell(row, columns[header]).value

            print()
            print(
                f"[{index}] Fila {row} | Fecha: {value('Fecha')} | "
                f"Referencia: {value('ReferenciaVenta') or '(sin referencia)'}"
            )
            print(
                f"    Origen: {value('Origen')} | "
                f"Ingreso: {value('Ingreso')} | Egreso: {value('Egreso')} | "
                f"Importe original: {value('Importe original')}"
            )
            print(f"    Concepto: {value('Concepto') or '(sin concepto)'}")
            print(f"    Observación: {value('Observación / ajuste') or '(sin observación)'}")
        print()
        print("Fin del detalle. El archivo de entrada no fue modificado.")
        return len(pending)
    finally:
        workbook.close()

def create_verified_backup(source: Path) -> Path:
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    backup = BACKUP_DIR / f"CarnePremiumERP_before_safe_apply_{stamp}.xlsx"
    shutil.copy2(source, backup)
    if sha256_file(source) != sha256_file(backup):
        backup.unlink(missing_ok=True)
        raise RuntimeError("El respaldo no coincide con el original; se cancela.")
    return backup


def run_build(source: Path, apply_result: bool) -> int:
    if not source.is_file():
        print(f"ERROR: no existe el archivo de entrada: {source}")
        return 2
    if source.suffix.lower() != ".xlsx" or not valid_xlsx(source):
        print("ERROR: el archivo de entrada no es un XLSX válido.")
        return 2

    source = source.resolve()
    lock_file = source.parent / ("~$" + source.name)
    if lock_file.exists():
        print("NO SE INICIÓ: Excel parece tener abierto el libro de entrada.")
        print(f"Cerrá Excel antes de continuar: {source}")
        return 2
    if apply_result and source != DEFAULT_MASTER.resolve():
        print("ERROR: --aplicar solo está permitido sobre el maestro operativo.")
        print("Para validar otros libros, ejecutá el actualizador sin --aplicar.")
        return 2

    original_hash = sha256_file(source)
    print("CARNE PREMIUM ERP — ACTUALIZADOR SEGURO")
    print(f"Entrada: {source}")
    print(f"Modo: {'VALIDAR Y APLICAR' if apply_result else 'SOLO VALIDAR'}")
    print(f"SHA-256 de entrada: {original_hash}")
    print("El constructor y las pruebas se ejecutan en una copia temporal.")

    with tempfile.TemporaryDirectory(prefix="CarnePremiumERP_build_") as temp:
        stage_root = Path(temp)
        stage_erp = stage_root / "01_ERP"
        stage_backups = stage_root / "02_Backups"
        stage_erp.mkdir(parents=True, exist_ok=True)
        stage_backups.mkdir(parents=True, exist_ok=True)
        staged_master = stage_erp / "CarnePremiumERP.xlsx"
        shutil.copy2(source, staged_master)

        # Los respaldos usados para recuperación histórica también quedan aislados.
        if BACKUP_DIR.exists():
            shutil.copytree(
                BACKUP_DIR,
                stage_backups,
                dirs_exist_ok=True,
                ignore=shutil.ignore_patterns("~$*", "*.tmp", "*.lock"),
            )

        env = os.environ.copy()
        env["CARNEPREMIUM_MASTER_FILE"] = str(staged_master)
        steps = [
            ("1/3 — Construcción de la copia temporal", SCRIPTS_DIR / "builder.py"),
            ("2/3 — Validación estructural y de fórmulas", TESTS_DIR / "test_mvp.py"),
            ("3/3 — Pruebas de Ticket, Caja y Cuenta Corriente", TESTS_DIR / "test_ticket_caja.py"),
        ]
        for label, script in steps:
            if not run_step(label, script, env):
                print("BUILD RECHAZADO. No se reemplazó el libro de entrada.")
                return 1

        state_source = stage_erp / STATE_NAME
        if not valid_xlsx(staged_master) or not state_source.is_file():
            print("ERROR: no se generó el XLSX temporal validado o su archivo de estado.")
            print("El libro de entrada no fue reemplazado.")
            return 1

        pending_adjustments = report_pending_adjustments(staged_master)

        new_hash = sha256_file(staged_master)
        if sha256_file(source) != original_hash:
            print("ERROR: el libro de entrada cambió durante la validación. No se aplicará.")
            return 3
        if not apply_result:
            print("VALIDACIÓN COMPLETADA: el libro de entrada no fue modificado.")
            print(f"SHA-256 de entrada: {sha256_file(source)}")
            print(f"SHA-256 del resultado validado: {new_hash}")
            return 0

        if lock_file.exists() or sha256_file(source) != original_hash:
            print("ERROR: el libro está abierto o cambió antes de aplicar. Se cancela.")
            return 3

        backup = create_verified_backup(source)
        state_target = source.parent / STATE_NAME
        previous_state = state_target.read_bytes() if state_target.exists() else None
        state_staging = stage_root / "state_to_apply.json"
        workbook_replaced = False
        try:
            state = json.loads(state_source.read_text(encoding="utf-8"))
            state["master_file"] = str(source)
            state["sha256"] = new_hash
            state["saved_at"] = datetime.now().isoformat(timespec="seconds")
            state_staging.write_text(
                json.dumps(state, indent=2, ensure_ascii=False),
                encoding="utf-8",
            )
            atomic_copy(staged_master, source)
            workbook_replaced = True
            atomic_copy(state_staging, state_target)
            if sha256_file(source) != new_hash:
                raise RuntimeError("La verificación SHA-256 posterior falló.")
        except Exception:
            if workbook_replaced:
                atomic_copy(backup, source)
            if previous_state is None:
                state_target.unlink(missing_ok=True)
            else:
                restore_state = stage_root / "state_to_restore.json"
                restore_state.write_bytes(previous_state)
                atomic_copy(restore_state, state_target)
            raise

        print("BUILD VALIDADO Y APLICADO.")
        print(f"Respaldo anterior verificado: {backup}")
        print(f"Archivo actualizado: {source}")
        print(f"SHA-256 aplicado: {new_hash}")
        return 0


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Valida Carne Premium ERP en una copia temporal. Por defecto no "
            "reemplaza ningún libro. --aplicar solo se permite sobre el maestro "
            "después de que todas las pruebas terminen correctamente."
        )
    )
    parser.add_argument(
        "--archivo",
        default=str(DEFAULT_MASTER),
        help="XLSX de entrada; por defecto el maestro operativo.",
    )
    parser.add_argument(
        "--aplicar",
        action="store_true",
        help="Aplica al maestro el resultado validado y crea un respaldo previo.",
    )
    args = parser.parse_args()
    source = Path(args.archivo).expanduser()
    if not source.is_absolute():
        source = PROJECT_ROOT / source
    try:
        return run_build(source.resolve(), args.aplicar)
    except Exception as exc:
        print(f"ERROR INESPERADO: {type(exc).__name__}: {exc}")
        print("No consideres aplicado el build si aparece este error.")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
