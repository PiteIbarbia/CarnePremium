from __future__ import annotations

from datetime import date, datetime, time
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_SALES_ROW = 501
CAJA_HEADERS = [
    "MovimientoID", "Fecha", "Tipo", "Concepto", "ReferenciaVenta",
    "MedioPago", "Ingreso", "Egreso", "Saldo", "Origen",
]
CAJA_AUDIT_HEADERS = ["Importe original", "Observación / ajuste"]


def _column_map(ws):
    result = {}
    for column in range(1, ws.max_column + 1):
        value = ws.cell(1, column).value
        if value not in (None, ""):
            result[str(value).strip()] = column
    return result


def _text(value):
    return str(value or "").strip()


def _number(value):
    if isinstance(value, bool) or value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        cleaned = value.strip().replace("$", "").replace(" ", "")
        if "," in cleaned and "." in cleaned:
            cleaned = cleaned.replace(".", "").replace(",", ".")
        elif "," in cleaned:
            cleaned = cleaned.replace(",", ".")
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def _sale_reference(ws, row, columns):
    value = ws.cell(row, columns["VentaID"]).value
    if isinstance(value, str) and value.startswith("="):
        return f"VEN{row - 1:06d}"
    if value in (None, ""):
        return f"VEN{row - 1:06d}"
    return str(value).strip()


def _column_has_data(ws, column):
    return any(
        ws.cell(row, column).value not in (None, "")
        for row in range(2, ws.max_row + 1)
    )


def _ensure_audit_columns(ws):
    """Añade columnas de revisión sin sobrescribir datos que ya existan."""
    columns = _column_map(ws)
    for header in CAJA_AUDIT_HEADERS:
        if header in columns:
            continue
        chosen = None
        upper_bound = max(ws.max_column + 12, len(CAJA_HEADERS) + 20)
        for column in range(len(CAJA_HEADERS) + 1, upper_bound + 1):
            current_header = _text(ws.cell(1, column).value)
            if current_header:
                continue
            if _column_has_data(ws, column):
                continue
            chosen = column
            break
        if chosen is None:
            raise RuntimeError(
                f"No se encontró una columna vacía para '{header}' en 07_CAJA. "
                "No se sobrescribieron los datos existentes."
            )
        ws.cell(1, chosen).value = header
        columns[header] = chosen
    return columns




def _prepare_cash_summary(ws, columns):
    """Crea un resumen lateral sin insertar filas ni sobrescribir movimientos."""
    title = "RESUMEN CAJA"
    summary_col = next(
        (
            column
            for column in range(1, ws.max_column + 1)
            if _text(ws.cell(1, column).value) == title
        ),
        None,
    )
    if summary_col is None:
        start_at = max(
            column
            for header, column in columns.items()
            if header in CAJA_HEADERS + CAJA_AUDIT_HEADERS
        ) + 2
        for candidate in range(start_at, start_at + 40):
            pair_is_empty = all(
                ws.cell(row, column).value in (None, "")
                for row in range(1, 8)
                for column in (candidate, candidate + 1)
            )
            if pair_is_empty:
                summary_col = candidate
                break
    if summary_col is None:
        raise RuntimeError(
            "No se encontró un espacio libre para el resumen de 07_CAJA. "
            "No se sobrescribieron datos."
        )

    summary_rows = {
        1: (title, None),
        2: ("Ingresos registrados", "=SUM(G2:G10000)"),
        3: ("Egresos registrados", "=SUM(H2:H10000)"),
        4: ("Saldo neto registrado", "=SUM(G2:G10000)-SUM(H2:H10000)"),
        5: ("Ajustes pendientes", '=COUNTIF(C2:C10000,"AJUSTE PENDIENTE")'),
        6: ("Saldo inicial de efectivo", "No incluido"),
        7: ("Nota", "Conciliá los ajustes antes de cerrar caja."),
    }
    label_letter = get_column_letter(summary_col)
    value_letter = get_column_letter(summary_col + 1)

    for row, (label, value) in summary_rows.items():
        label_cell = ws.cell(row, summary_col)
        value_cell = ws.cell(row, summary_col + 1)
        label_cell.value = label
        value_cell.value = value
        label_cell.alignment = Alignment(vertical="center", wrap_text=True)
        value_cell.alignment = Alignment(vertical="center", wrap_text=True)

        if row == 1:
            for cell in (label_cell, value_cell):
                cell.fill = PatternFill("solid", fgColor="7A3E13")
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
        else:
            label_cell.font = Font(bold=True, color="5B341B")
            label_cell.fill = PatternFill("solid", fgColor="F4EDE7")
            value_cell.fill = PatternFill("solid", fgColor="FBF8F5")
            if row in (2, 3, 4):
                value_cell.number_format = '$ #,##0.00;[Red]-$ #,##0.00'
    ws.column_dimensions[label_letter].width = max(
        ws.column_dimensions[label_letter].width or 0, 25
    )
    ws.column_dimensions[value_letter].width = max(
        ws.column_dimensions[value_letter].width or 0, 45
    )
    ws.row_dimensions[1].height = max(ws.row_dimensions[1].height or 0, 30)
    ws.row_dimensions[7].height = max(ws.row_dimensions[7].height or 0, 34)


def preparar_caja(workbook):
    """Prepara Caja conservando los movimientos e incorporando revisión auditable."""
    ws = workbook["07_CAJA"] if "07_CAJA" in workbook.sheetnames else workbook.create_sheet("07_CAJA")
    current = [_text(ws.cell(1, c).value) for c in range(1, len(CAJA_HEADERS) + 1)]
    has_rows = any(
        ws.cell(row, column).value not in (None, "")
        for row in range(2, ws.max_row + 1)
        for column in range(1, len(CAJA_HEADERS) + 1)
    )
    if all(value == "" for value in current) and not has_rows:
        for column, header in enumerate(CAJA_HEADERS, start=1):
            ws.cell(1, column).value = header
    elif current != CAJA_HEADERS:
        raise RuntimeError(
            "La estructura de 07_CAJA no coincide con el esquema de este Build. "
            "No se modificó para proteger los movimientos existentes."
        )

    columns = _ensure_audit_columns(ws)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.tabColor = "7A3E13"
    ws.freeze_panes = "D2"

    widths = {
        "A": 16, "B": 19, "C": 23, "D": 42, "E": 20,
        "F": 20, "G": 17, "H": 17, "I": 19, "J": 23,
    }
    for column, header in enumerate(CAJA_HEADERS, start=1):
        cell = ws.cell(1, column)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7A3E13")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for header, width in zip(CAJA_AUDIT_HEADERS, (18, 54)):
        column = columns[header]
        ws.cell(1, column).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, column).fill = PatternFill("solid", fgColor="7A3E13")
        ws.cell(1, column).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        widths[get_column_letter(column)] = width

    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.row_dimensions[1].height = 34
    header_comments = {
        "G": "Dinero realmente recibido. No uses esta columna para registrar el total vendido si todavía no se cobró.",
        "H": "Dinero realmente entregado o pagado. Las anulaciones técnicas automáticas no son egresos reales.",
        "I": "Saldo acumulado de los movimientos ordenados por fecha: ingresos menos egresos. No incluye un saldo inicial de efectivo.",
    }
    for address, text in header_comments.items():
        header_cell = ws[f"{address}1"]
        if header_cell.comment is None:
            header_cell.comment = Comment(text, "Carne Premium ERP")
    for header, text in {
        "Importe original": "Importe conservado solo como referencia cuando un movimiento automático queda pendiente de revisión. No suma ni resta al saldo.",
        "Observación / ajuste": "Explica por qué el movimiento necesita revisión. Una marca de revisión no acredita por sí sola una entrada o salida real de dinero.",
    }.items():
        column = columns[header]
        if ws.cell(1, column).comment is None:
            ws.cell(1, column).comment = Comment(text, "Carne Premium ERP")

    _prepare_cash_summary(ws, columns)
    return ws


def _next_movement_id(ws, movement_col):
    largest = 0
    for row in range(2, ws.max_row + 1):
        value = _text(ws.cell(row, movement_col).value)
        if value.startswith("MOV"):
            try:
                largest = max(largest, int(value[3:]))
            except ValueError:
                pass
    return f"MOV{largest + 1:06d}"


def _summary_columns(ws):
    """Devuelve las columnas fijas del resumen lateral, si existe."""
    for column in range(1, ws.max_column + 1):
        if _text(ws.cell(1, column).value) == "RESUMEN CAJA":
            return {column, column + 1}
    return set()


def _last_movement_row(ws, columns):
    """Busca el último movimiento ignorando el saldo calculado y el resumen lateral."""
    ignored = {columns.get("Saldo")} | _summary_columns(ws)
    last = 1
    for row in range(2, ws.max_row + 1):
        if any(
            ws.cell(row, column).value not in (None, "")
            for column in range(1, ws.max_column + 1)
            if column not in ignored
        ):
            last = row
    return last


def _client_id_for_name(workbook, client_name):
    target = _text(client_name).casefold()
    if not target or "03_CLIENTES" not in workbook.sheetnames:
        return ""
    clients = workbook["03_CLIENTES"]
    for row in range(2, min(MAX_SALES_ROW, clients.max_row) + 1):
        name = _text(clients.cell(row, 2).value)
        if not name or name.casefold() != target:
            continue
        raw_id = clients.cell(row, 1).value
        if raw_id in (None, "") or (isinstance(raw_id, str) and raw_id.startswith("=")):
            return f"CLI{row - 1:06d}"
        return str(raw_id).strip()
    return ""


def _cash_date(value, fallback=None):
    if isinstance(value, (date, datetime)):
        return value
    if isinstance(fallback, (date, datetime)):
        return fallback
    return datetime.now()


def _date_sort_value(value):
    if isinstance(value, datetime):
        return (0, value)
    if isinstance(value, date):
        return (0, datetime.combine(value, time.min))
    if isinstance(value, str):
        raw = value.strip()
        for parser in (datetime.fromisoformat,):
            try:
                return (0, parser(raw))
            except ValueError:
                pass
        for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d"):
            try:
                return (0, datetime.strptime(raw, fmt))
            except ValueError:
                pass
    return (1, datetime.max)


def _mark_for_review(ws, columns, row, reason):
    """Neutraliza un ajuste automático dudoso y conserva su importe como referencia."""
    income = _number(ws.cell(row, columns["Ingreso"]).value) or 0.0
    expense = _number(ws.cell(row, columns["Egreso"]).value) or 0.0
    original_column = columns.get("Importe original")
    note_column = columns.get("Observación / ajuste")

    if original_column:
        existing_original = _number(ws.cell(row, original_column).value) or 0.0
        original = existing_original if existing_original > 0 else max(abs(income), abs(expense))
        if original > 0:
            ws.cell(row, original_column).value = round(original, 2)

    ws.cell(row, columns["Tipo"]).value = "AJUSTE PENDIENTE"
    ws.cell(row, columns["Concepto"]).value = reason
    # Nunca convertir una anulación técnica en una salida de efectivo ficticia.
    ws.cell(row, columns["Ingreso"]).value = 0
    ws.cell(row, columns["Egreso"]).value = 0
    ws.cell(row, columns["Origen"]).value = "VENTA_AUTO_REVISION"
    if note_column:
        ws.cell(row, note_column).value = (
            "Revisar contra 05_VENTAS y, si corresponde, las copias de respaldo. "
            f"Motivo: {reason}. "
            "Este ajuste queda fuera del saldo hasta confirmar el movimiento real."
        )


def _sort_cash_rows(ws, columns):
    """Ordena movimientos por fecha, conservando juntos sus datos y el resumen lateral."""
    original_max_row = ws.max_row
    last_column = ws.max_column
    saldo_column = columns["Saldo"]
    fixed_columns = _summary_columns(ws)
    records = []

    for row in range(2, original_max_row + 1):
        values = [ws.cell(row, column).value for column in range(1, last_column + 1)]
        # El saldo es fórmula derivada y el resumen lateral no pertenece a los movimientos.
        has_movement = any(
            value not in (None, "")
            for column, value in enumerate(values, start=1)
            if column != saldo_column and column not in fixed_columns
        )
        if has_movement:
            date_value = values[columns["Fecha"] - 1]
            records.append((_date_sort_value(date_value), row, values))

    records.sort(key=lambda item: (item[0][0], item[0][1], item[1]))

    # Vaciar solamente las columnas de movimientos. Las columnas del resumen se
    # quedan quietas; cualquier columna operativa adicional viaja con su movimiento.
    for row in range(2, original_max_row + 1):
        for column in range(1, last_column + 1):
            if column not in fixed_columns:
                ws.cell(row, column).value = None

    for target_row, (_, _, values) in enumerate(records, start=2):
        for column, value in enumerate(values, start=1):
            if column not in fixed_columns:
                ws.cell(target_row, column).value = value

    return len(records) + 1


def _apply_cash_presentation(ws, columns, last_row):
    last_header_col = max(columns.values())
    ws.auto_filter.ref = f"A1:{get_column_letter(last_header_col)}{max(1, last_row)}"

    for row in range(2, last_row + 1):
        ws.cell(row, columns["Fecha"]).number_format = "dd/mm/yyyy hh:mm"
        for header in ("Ingreso", "Egreso", "Saldo", "Importe original"):
            column = columns.get(header)
            if column:
                ws.cell(row, column).number_format = '$ #,##0.00;[Red]-$ #,##0.00'

    # Resaltar por el tipo de movimiento con formato condicional. Así el color
    # acompaña al movimiento cuando se ordena la hoja en actualizaciones posteriores.
    has_pending_rule = False
    for conditional_range in ws.conditional_formatting:
        for rule in ws.conditional_formatting[conditional_range]:
            formula = [str(item) for item in (getattr(rule, "formula", []) or [])]
            if any("AJUSTE PENDIENTE" in item for item in formula):
                has_pending_rule = True
                break
        if has_pending_rule:
            break
    if not has_pending_rule:
        ws.conditional_formatting.add(
            f"A2:{get_column_letter(last_header_col)}10000",
            FormulaRule(
                formula=['$C2="AJUSTE PENDIENTE"'],
                fill=PatternFill("solid", fgColor="FFF2CC"),
                font=Font(color="7F6000"),
            ),
        )

    # No se eliminan reglas ajenas; se evita agregar una regla idéntica en cada Build.
    has_negative_rule = False
    for conditional_range in ws.conditional_formatting:
        for rule in ws.conditional_formatting[conditional_range]:
            if (
                getattr(rule, "type", None) == "cellIs"
                and getattr(rule, "operator", None) == "lessThan"
                and "0" in (getattr(rule, "formula", []) or [])
                and "I" in str(conditional_range.sqref)
            ):
                has_negative_rule = True
                break
        if has_negative_rule:
            break
    if not has_negative_rule:
        ws.conditional_formatting.add(
            "I2:I10000",
            CellIsRule(
                operator="lessThan",
                formula=["0"],
                fill=PatternFill("solid", fgColor="FCE4D6"),
                font=Font(color="9C0006", bold=True),
            ),
        )


def sincronizar_ventas_en_caja(workbook):
    """
    Registra solo dinero efectivamente cobrado.
    Las anomalías automáticas quedan marcadas y auditables, nunca como egresos ficticios.
    El saldo se recalcula después de ordenar los movimientos por fecha.
    """
    cash = preparar_caja(workbook)
    sales = workbook["05_VENTAS"]
    sc = _column_map(sales)
    cc = {
        name: column
        for name, column in _column_map(cash).items()
        if name in CAJA_HEADERS + CAJA_AUDIT_HEADERS
    }
    required = [
        "VentaID", "Fecha", "Cliente", "Producto", "Importe",
        "PagoRecibido", "MedioPago", "EstadoVenta", "TipoRegistro",
    ]
    missing = [name for name in required if name not in sc]
    if missing:
        raise RuntimeError("05_VENTAS: faltan columnas para Caja: " + ", ".join(missing))

    # Migrar anulaciones técnicas creadas por versiones previas. Se conserva el
    # rastro y el importe original, pero no se considera una salida real de dinero.
    for row in range(2, cash.max_row + 1):
        origin = _text(cash.cell(row, cc["Origen"]).value).upper()
        kind = _text(cash.cell(row, cc["Tipo"]).value).upper()
        if origin == "VENTA_AUTO_ANULADA" or (
            origin.startswith("VENTA_AUTO") and kind in {"ANULACION", "REVISION COBRO"}
        ):
            ref = _text(cash.cell(row, cc["ReferenciaVenta"]).value)
            _mark_for_review(
                cash, cc, row,
                f"Anulación automática heredada pendiente de revisar: {ref or 'sin referencia'}",
            )

    expected = {}
    known_refs = set()
    for row in range(2, min(MAX_SALES_ROW, sales.max_row) + 1):
        status = _text(sales.cell(row, sc["EstadoVenta"]).value).upper()
        kind = _text(sales.cell(row, sc["TipoRegistro"]).value).upper() or "VENTA"
        if status == "OK" and kind != "COBRO":
            ref = _sale_reference(sales, row, sc)
            known_refs.add(ref)
        elif status == "COBRO OK" and kind == "COBRO":
            ref = f"COB{row - 1:06d}"
            known_refs.add(ref)
        else:
            continue

        client = _text(sales.cell(row, sc["Cliente"]).value)
        client_id = _client_id_for_name(workbook, client)
        paid = _number(sales.cell(row, sc["PagoRecibido"]).value)
        payment = _text(sales.cell(row, sc["MedioPago"]).value)
        sale_date = _cash_date(sales.cell(row, sc["Fecha"]).value)
        payment_date_col = sc.get("FechaCobro")
        payment_date_value = sales.cell(row, payment_date_col).value if payment_date_col else None
        collection_date = _cash_date(payment_date_value, fallback=sale_date)

        if (
            not client or not client_id or paid is None or paid <= 0
            or not payment
            or payment.casefold().replace(" ", "") == "cuentacorriente"
        ):
            continue

        if kind == "COBRO" and status == "COBRO OK":
            concept = f"Cobro de saldo anterior - {client}"
            ref = f"COB{row - 1:06d}"
        else:
            product = _text(sales.cell(row, sc["Producto"]).value)
            if not product:
                continue
            ticket_id = f"T-{sale_date:%Y%m%d}-{client_id}"
            concept = f"Cobro {ticket_id} - {client} / {product}"
            ref = _sale_reference(sales, row, sc)

        expected[ref] = {
            "fecha": collection_date,
            "concepto": concept,
            "medio_pago": payment,
            "importe": round(paid, 2),
        }

    # Los duplicados automáticos se conservan como revisión, en lugar de borrarse.
    duplicates = []
    seen = set()
    for row in range(2, cash.max_row + 1):
        if _text(cash.cell(row, cc["Origen"]).value) != "VENTA_AUTO":
            continue
        ref = _text(cash.cell(row, cc["ReferenciaVenta"]).value)
        if ref and ref in seen:
            duplicates.append((row, ref))
        elif ref:
            seen.add(ref)
    for row, ref in duplicates:
        _mark_for_review(
            cash, cc, row,
            f"Movimiento automático duplicado pendiente de revisar: {ref}",
        )

    existing = {}
    for row in range(2, cash.max_row + 1):
        if _text(cash.cell(row, cc["Origen"]).value) == "VENTA_AUTO":
            ref = _text(cash.cell(row, cc["ReferenciaVenta"]).value)
            if ref:
                existing[ref] = row

    for ref, row in list(existing.items()):
        if ref in expected:
            continue
        if ref in known_refs:
            reason = f"Cobro previo sin confirmación en 05_VENTAS: {ref}"
        else:
            reason = f"Cobro automático sin venta asociada identificable: {ref}"
        _mark_for_review(cash, cc, row, reason)

    active = {}
    for row in range(2, cash.max_row + 1):
        if _text(cash.cell(row, cc["Origen"]).value) == "VENTA_AUTO":
            ref = _text(cash.cell(row, cc["ReferenciaVenta"]).value)
            if ref:
                active[ref] = row

    added = updated = 0
    for ref, receipt in expected.items():
        row = active.get(ref)
        if row is None:
            row = _last_movement_row(cash, cc) + 1
            cash.cell(row, cc["MovimientoID"]).value = _next_movement_id(cash, cc["MovimientoID"])
            added += 1
        else:
            updated += 1
        values = {
            "Fecha": receipt["fecha"],
            "Tipo": "INGRESO",
            "Concepto": receipt["concepto"],
            "ReferenciaVenta": ref,
            "MedioPago": receipt["medio_pago"],
            "Ingreso": receipt["importe"],
            "Egreso": 0,
            "Origen": "VENTA_AUTO",
        }
        for header, value in values.items():
            cash.cell(row, cc[header]).value = value

    last_row = _sort_cash_rows(cash, cc)
    for row in range(2, last_row + 1):
        cash.cell(row, cc["Saldo"]).value = f"=SUM($G$2:G{row})-SUM($H$2:H{row})"
    _apply_cash_presentation(cash, cc, last_row)

    return {
        "agregados": added,
        "actualizados": updated,
        "movimientos": max(0, last_row - 1),
        "ingresos_cobrados": round(sum(item["importe"] for item in expected.values()), 2),
        "ajustes_pendientes": sum(
            1 for row in range(2, last_row + 1)
            if _text(cash.cell(row, cc["Tipo"]).value).upper() == "AJUSTE PENDIENTE"
        ),
    }
