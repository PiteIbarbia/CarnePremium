from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from pathlib import Path
import warnings

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
MASTER = ROOT / '01_ERP' / 'CarnePremiumERP.xlsx'
MAX_ROW = 501

def txt(value):
    return str(value or '').strip()

def num(value):
    if isinstance(value, bool) or value in (None, ''):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        raw = value.strip().replace('$', '').replace(' ', '')
        if ',' in raw and '.' in raw:
            raw = raw.replace('.', '').replace(',', '.')
        elif ',' in raw:
            raw = raw.replace(',', '.')
        try:
            return float(raw)
        except ValueError:
            return None
    return None

def is_formula(value):
    return isinstance(value, str) and value.startswith('=')

def main():
    print('=' * 64)
    print('DIAGNÓSTICO LOCAL DEL ERP — SOLO LECTURA')
    print('No modifica el Excel y no muestra nombres ni importes.')
    print('=' * 64)
    if not MASTER.exists():
        print('ERROR: no se encuentra 01_ERP/CarnePremiumERP.xlsx.')
        return 2
    lock = MASTER.parent / ('~$' + MASTER.name)
    print(f'Archivo maestro: OK | Excel abierto: {"SÍ" if lock.exists() else "NO"}')
    if lock.exists():
        print('AVISO: para una comprobación fiable, guardá y cerrá Excel antes de repetir.')

    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        wb = load_workbook(MASTER, data_only=False)
        wbv = load_workbook(MASTER, data_only=True)
    try:
        missing_sheets = [s for s in ('05_VENTAS', '06_TICKET') if s not in wb.sheetnames]
        if missing_sheets:
            print('ERROR: faltan hojas: ' + ', '.join(missing_sheets))
            return 2
        sales = wb['05_VENTAS']
        headers = {txt(sales.cell(1, c).value): c for c in range(1, sales.max_column + 1) if sales.cell(1, c).value not in (None, '')}
        required = ('Fecha', 'Cliente', 'Producto', 'Kg', 'PrecioKg', 'Importe', 'Ganancia$', 'Ganancia%', 'EstadoVenta', 'PagoRecibido', '_TicketLineaKey', 'TipoRegistro')
        missing_headers = [name for name in required if name not in headers]
        print()
        print('[1] 05_VENTAS')
        print(f'Encabezados requeridos: {"OK" if not missing_headers else "FALTAN: " + ", ".join(missing_headers)}')
        if missing_headers:
            return 2
        statuses, kinds = Counter(), Counter()
        formulas, literals, blanks = (Counter() for _ in range(3))
        eligible = missing_key = valid_dates = row_count = 0
        for row in range(2, MAX_ROW + 1):
            client = txt(sales.cell(row, headers['Cliente']).value)
            product = txt(sales.cell(row, headers['Producto']).value)
            kg = num(sales.cell(row, headers['Kg']).value)
            payment = num(sales.cell(row, headers['PagoRecibido']).value)
            kind = txt(sales.cell(row, headers['TipoRegistro']).value).upper()
            status = txt(sales.cell(row, headers['EstadoVenta']).value).upper()
            if not (client or product or kg is not None or payment is not None or kind):
                continue
            row_count += 1
            statuses[status or '(vacío)'] += 1
            kinds[kind or 'VENTA (vacío)'] += 1
            for letter in ('H', 'I', 'J', 'K'):
                value = sales[f'{letter}{row}'].value
                if is_formula(value):
                    formulas[letter] += 1
                elif value in (None, ''):
                    blanks[letter] += 1
                else:
                    literals[letter] += 1
            day = sales.cell(row, headers['Fecha']).value
            if isinstance(day, (date, datetime)):
                valid_dates += 1
            if status == 'OK' and kind != 'COBRO' and client and product and kg is not None and kg > 0 and isinstance(day, (date, datetime)):
                eligible += 1
                key = txt(sales.cell(row, headers['_TicketLineaKey']).value)
                if not (key.startswith('T-') and '|' in key):
                    missing_key += 1
        print(f'Filas con datos: {row_count}')
        print(f'Estados (conteos): {dict(statuses)}')
        print(f'Tipos de registro (conteos): {dict(kinds)}')
        print(f'Filas con fecha guardada como fecha: {valid_dates}')
        print('H:K — fórmulas / valores fijos / vacías:')
        for letter in ('H', 'I', 'J', 'K'):
            print(f'  {letter}: {formulas[letter]} / {literals[letter]} / {blanks[letter]}')

        print()
        print('[2] 06_TICKET')
        ticket, ticketv = wb['06_TICKET'], wbv['06_TICKET']
        selected = txt(ticket['B3'].value)
        listed_ids = {txt(ticket.cell(row, 7).value) for row in range(2, MAX_ROW + 1) if txt(ticket.cell(row, 7).value)}
        keys = [txt(sales.cell(row, headers['_TicketLineaKey']).value) for row in range(2, MAX_ROW + 1)]
        matching = sum(1 for key in keys if selected and key.startswith(selected + '|'))
        detail_formulas = sum(1 for row in range(9, 39) for col in range(1, 5) if is_formula(ticket.cell(row, col).value))
        detail_cached = sum(1 for row in range(9, 39) for col in range(1, 5) if ticketv.cell(row, col).value not in (None, ''))
        header_formulas = sum(1 for address in ('B4', 'B5', 'B6', 'B7') if is_formula(ticket[address].value))
        header_cached = sum(1 for address in ('B4', 'B5', 'B6', 'B7') if ticketv[address].value not in (None, ''))
        print(f'Tickets en lista: {len(listed_ids)}')
        print(f'Selector B3 con valor: {"SÍ" if selected else "NO"}')
        print(f'Selector B3 incluido en la lista: {"SÍ" if selected and selected in listed_ids else "NO"}')
        print(f'Ventas aptas para ticket: {eligible}')
        print(f'Ventas aptas sin clave X: {missing_key}')
        print(f'Claves que coinciden con selector actual: {matching}')
        print(f'Fórmulas de detalle: {detail_formulas}/120')
        print(f'Celdas de detalle con valor en caché: {detail_cached}/120')
        print(f'Fórmulas de encabezado: {header_formulas}/4')
        print(f'Encabezados con valor en caché: {header_cached}/4')

        print()
        print('[3] Hallazgos')
        issues = []
        if any(literals[column] for column in ('H', 'I', 'J', 'K')):
            issues.append('H:K contiene valores fijos en filas con datos; no todas conservan fórmulas.')
        if eligible and not listed_ids:
            issues.append('Hay ventas aptas para ticket, pero la lista de tickets está vacía.')
        if selected and selected not in listed_ids:
            issues.append('El selector actual no corresponde a un ticket listado.')
        if selected and matching == 0:
            issues.append('El selector actual no encuentra claves de línea en Ventas.')
        if missing_key:
            issues.append('Hay ventas aptas sin clave materializada de ticket.')
        if detail_formulas != 120 or header_formulas != 4:
            issues.append('Faltan fórmulas en la estructura de 06_TICKET.')
        for issue in issues:
            print('- ' + issue)
        if not issues:
            print('No se detectaron anomalías estructurales en los campos revisados.')
        print('Fin. El Excel no fue modificado.')
        return 0
    finally:
        wb.close()
        wbv.close()

if __name__ == '__main__':
    raise SystemExit(main())
